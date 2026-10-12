"""Small, source-locked Docling conversion diagnostic for synthetic office/web/email files.

This measures format conversion only. It does not run product extraction or an AI model.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import os
import platform
import subprocess
import sys
import threading
import time
from datetime import datetime, timezone
from email import policy
from email.message import EmailMessage
from importlib import metadata
from pathlib import Path
from typing import Any


FIXTURE_NAMES = ("order.docx", "order.xlsx", "order.html", "order.eml", "order.pptx")
HEADERS = ("Item", "Qty", "Amount")
ROWS = (("Widget A", "2", "12.00"), ("Widget B", "1", "7.50"))
FIXED_TIME = datetime(2020, 1, 1, tzinfo=timezone.utc)


def sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def html_body() -> str:
    return """<!doctype html>
<html lang="en"><head><meta charset="utf-8"><title>Synthetic order</title></head>
<body><h1>Synthetic order SYN-1001</h1><p>Order date: 2020-01-01</p>
<table><thead><tr><th>Item</th><th>Qty</th><th>Amount</th></tr></thead>
<tbody><tr><td>Widget A</td><td>2</td><td>12.00</td></tr>
<tr><td>Widget B</td><td>1</td><td>7.50</td></tr></tbody></table></body></html>
"""


def generate_fixtures(directory: Path) -> dict[str, Any]:
    from docx import Document
    from openpyxl import Workbook
    from openpyxl.styles import Font
    from openpyxl.worksheet.table import Table, TableStyleInfo
    from pptx import Presentation
    from pptx.util import Inches

    directory.mkdir(parents=True, exist_ok=True)
    if any((directory / name).exists() for name in (*FIXTURE_NAMES, "manifest.json")):
        raise FileExistsError("Use an empty directory so frozen fixture hashes are not replaced")

    docx = Document()
    docx.core_properties.title = "Synthetic order SYN-1001"
    docx.core_properties.created = FIXED_TIME
    docx.core_properties.modified = FIXED_TIME
    docx.add_heading("Synthetic order SYN-1001", level=1)
    docx.add_paragraph("Order date: 2020-01-01")
    table = docx.add_table(rows=1, cols=3)
    for cell, value in zip(table.rows[0].cells, HEADERS):
        cell.text = value
    for row_values in ROWS:
        for cell, value in zip(table.add_row().cells, row_values):
            cell.text = value
    docx.save(directory / "order.docx")

    workbook = Workbook()
    workbook.properties.title = "Synthetic order SYN-1001"
    workbook.properties.created = FIXED_TIME
    workbook.properties.modified = FIXED_TIME
    sheet = workbook.active
    sheet.title = "Order Lines"
    sheet.append(HEADERS)
    for row_values in ROWS:
        sheet.append(row_values)
    for cell in sheet[1]:
        cell.font = Font(bold=True)
    excel_table = Table(displayName="OrderLines", ref="A1:C3")
    excel_table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
    sheet.add_table(excel_table)
    workbook.save(directory / "order.xlsx")

    (directory / "order.html").write_text(html_body(), encoding="utf-8", newline="\n")

    message = EmailMessage()
    message["From"] = "sender@example.invalid"
    message["To"] = "recipient@example.invalid"
    message["Subject"] = "Synthetic order SYN-1001"
    message["Date"] = "Wed, 01 Jan 2020 00:00:00 +0000"
    message["Message-ID"] = "<syn-1001@example.invalid>"
    message.set_content("Synthetic order SYN-1001\nOrder date: 2020-01-01\n")
    message.add_alternative(html_body(), subtype="html")
    message.set_boundary("docai-synthetic-boundary-v1")
    (directory / "order.eml").write_bytes(message.as_bytes(policy=policy.SMTP))

    presentation = Presentation()
    presentation.core_properties.title = "Synthetic order SYN-1001"
    presentation.core_properties.created = FIXED_TIME
    presentation.core_properties.modified = FIXED_TIME
    slide = presentation.slides.add_slide(presentation.slide_layouts[5])
    slide.shapes.title.text = "Synthetic order SYN-1001"
    ppt_table = slide.shapes.add_table(3, 3, Inches(1), Inches(1.5), Inches(8), Inches(2)).table
    for row_index, values in enumerate((HEADERS, *ROWS)):
        for column_index, value in enumerate(values):
            ppt_table.cell(row_index, column_index).text = value
    presentation.save(directory / "order.pptx")

    manifest = {
        "fixture_set": "mixed_format_docling_v1",
        "synthetic": True,
        "expected_table_rows_per_source": 3,
        "expected_table_cells_per_source": 9,
        "files": {name: sha256(directory / name) for name in FIXTURE_NAMES},
    }
    (directory / "manifest.json").write_text(json.dumps(manifest, indent=2) + "\n", encoding="utf-8")
    return manifest


def verify_fixtures(directory: Path) -> dict[str, Any]:
    manifest = json.loads((directory / "manifest.json").read_text(encoding="utf-8"))
    if set(manifest["files"]) != set(FIXTURE_NAMES):
        raise ValueError("Fixture manifest has an unexpected file set")
    for name, expected in manifest["files"].items():
        actual = sha256(directory / name)
        if actual != expected:
            raise ValueError(f"Source hash mismatch: {name}")
    return manifest


def summarize_document(document: Any) -> dict[str, Any]:
    tables = []
    for table in document.tables:
        data = table.data
        cells = list(getattr(data, "table_cells", []) or [])
        tables.append({
            "rows": getattr(data, "num_rows", None),
            "columns": getattr(data, "num_cols", None),
            "cell_count": len(cells),
            "cell_texts": [str(getattr(cell, "text", "")) for cell in cells],
            "cell_grid": [{
                "row_start": cell.start_row_offset_idx,
                "row_end": cell.end_row_offset_idx,
                "column_start": cell.start_col_offset_idx,
                "column_end": cell.end_col_offset_idx,
                "row_span": cell.row_span,
                "column_span": cell.col_span,
            } for cell in cells],
            "table_page_numbers": sorted({prov.page_no for prov in (getattr(table, "prov", None) or []) if prov.page_no is not None}),
            "table_bbox_count": sum(prov.bbox is not None for prov in (getattr(table, "prov", None) or [])),
            "table_bboxes": [prov.bbox.model_dump() for prov in (getattr(table, "prov", None) or []) if prov.bbox is not None],
            "cell_bbox_count": sum(getattr(cell, "bbox", None) is not None for cell in cells),
        })

    item_count = 0
    item_provenance_count = 0
    item_bbox_count = 0
    page_numbers: set[int] = set()
    for item, _ in document.iterate_items():
        item_count += 1
        provenance = getattr(item, "prov", None) or []
        if provenance:
            item_provenance_count += 1
        for prov in provenance:
            if prov.page_no is not None:
                page_numbers.add(prov.page_no)
            if prov.bbox is not None:
                item_bbox_count += 1

    markdown = document.export_to_markdown()
    groups = [{"label": str(group.label), "name": group.name} for group in document.groups]
    return {
        "document_name": document.name,
        "pages_count": len(document.pages),
        "page_keys": sorted(document.pages),
        "page_sizes": {str(number): page.size.model_dump() for number, page in document.pages.items()},
        "groups": groups,
        "sheet_group_names": [group["name"] for group in groups if group["label"].endswith("sheet")],
        "item_count": item_count,
        "items_with_provenance": item_provenance_count,
        "item_bbox_count": item_bbox_count,
        "provenance_page_numbers": sorted(page_numbers),
        "table_count": len(tables),
        "table_rows_total": sum(table["rows"] or 0 for table in tables),
        "table_cells_total": sum(table["cell_count"] for table in tables),
        "tables": tables,
        "identifier_seen_in_markdown": "SYN-1001" in markdown,
        "sheet_name_seen_in_markdown": "Order Lines" in markdown,
        "markdown_character_count": len(markdown),
        "markdown_excerpt": markdown[:500],
    }


def convert_one(path: Path) -> dict[str, Any]:
    import psutil

    process = psutil.Process()
    baseline_rss = process.memory_info().rss
    peak_rss = baseline_rss
    stop = threading.Event()

    def sample_memory() -> None:
        nonlocal peak_rss
        while not stop.wait(0.01):
            peak_rss = max(peak_rss, process.memory_info().rss)

    sampler = threading.Thread(target=sample_memory, daemon=True)
    sampler.start()
    started = time.perf_counter()
    output: dict[str, Any] = {"source": path.name, "source_sha256": sha256(path)}
    try:
        from docling.datamodel.base_models import InputFormat
        from docling.document_converter import DocumentConverter

        format_by_suffix = {
            ".docx": InputFormat.DOCX,
            ".xlsx": InputFormat.XLSX,
            ".html": InputFormat.HTML,
            ".eml": InputFormat.EMAIL,
            ".pptx": InputFormat.PPTX,
        }
        converter = DocumentConverter(allowed_formats=[format_by_suffix[path.suffix.lower()]])
        converted_at = time.perf_counter()
        result = converter.convert(path, raises_on_error=False)
        finished_at = time.perf_counter()
        output.update({
            "status": str(result.status),
            "init_seconds": round(converted_at - started, 4),
            "conversion_seconds": round(finished_at - converted_at, 4),
            "errors": [str(error)[:300] for error in result.errors],
        })
        if result.document is not None:
            output.update(summarize_document(result.document))
    except Exception as exc:
        output.update({"status": "exception", "error": f"{type(exc).__name__}: {exc}"[:500]})
    finally:
        stop.set()
        sampler.join(timeout=1)
        peak_rss = max(peak_rss, process.memory_info().rss)
        output["peak_rss_mib"] = round(peak_rss / (1024 * 1024), 2)
        output["baseline_rss_mib"] = round(baseline_rss / (1024 * 1024), 2)
        output["worker_wall_seconds"] = round(time.perf_counter() - started, 4)
    return output


def evaluate(directory: Path, output_path: Path) -> dict[str, Any]:
    manifest = verify_fixtures(directory)
    records = []
    for name in FIXTURE_NAMES:
        started = time.perf_counter()
        environment = dict(os.environ, HF_HUB_OFFLINE="1", TRANSFORMERS_OFFLINE="1")
        proc = subprocess.run(
            [sys.executable, __file__, "worker", str(directory / name)],
            capture_output=True,
            text=True,
            timeout=120,
            env=environment,
            check=False,
        )
        try:
            record = json.loads(proc.stdout.strip())
        except json.JSONDecodeError:
            record = {"source": name, "status": "worker_failed", "error": proc.stdout[-500:]}
        expected_cells = [*HEADERS, *(cell for row in ROWS for cell in row)]
        actual_cells = [cell for table in record.get("tables", []) for cell in table["cell_texts"]]
        record["expected_table_cells"] = len(expected_cells)
        record["exact_table_cell_matches"] = sum(a == b for a, b in zip(expected_cells, actual_cells))
        record["process_wall_seconds"] = round(time.perf_counter() - started, 4)
        record["exit_code"] = proc.returncode
        record["stderr_excerpt"] = proc.stderr[-500:]
        records.append(record)
        print(f"{name}: {record['status']} in {record['process_wall_seconds']} s", file=sys.stderr, flush=True)

    package_names = ("docling-slim", "docling-core", "python-docx", "openpyxl", "python-pptx", "mail-parser", "beautifulsoup4", "psutil")
    result = {
        "protocol": "mixed_format_docling_v1",
        "fixture_manifest": manifest,
        "runtime": {"python": sys.version.split()[0], "platform": platform.platform(),
                    "packages": {name: metadata.version(name) for name in package_names}},
        "gpu_used": False,
        "model_downloads": False,
        "records": records,
    }
    output_path.parent.mkdir(parents=True, exist_ok=True)
    output_path.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    generate = commands.add_parser("generate")
    generate.add_argument("directory", type=Path)
    run = commands.add_parser("evaluate")
    run.add_argument("directory", type=Path)
    run.add_argument("output", type=Path)
    worker = commands.add_parser("worker")
    worker.add_argument("source", type=Path)
    args = parser.parse_args()

    if args.command == "generate":
        print(json.dumps(generate_fixtures(args.directory), indent=2))
    elif args.command == "evaluate":
        evaluate(args.directory, args.output)
    else:
        print(json.dumps(convert_one(args.source), ensure_ascii=False))


if __name__ == "__main__":
    main()
