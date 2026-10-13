"""Parser-only, source-locked diagnostic for the real public native dev corpus.

No extraction model, product workflow, network call, source text, or gold label is
used or written. Run only after the frozen protocol authorizes local parsing.
"""

from __future__ import annotations

import argparse
import collections
import hashlib
import json
import os
import platform
import re
import subprocess
import sys
import threading
import time
from importlib import metadata
from pathlib import Path
from typing import Any

MANIFEST = Path(__file__).resolve().parents[1] / "docs/evaluations/NATIVE_REAL_PUBLIC_DEV_MANIFEST_20261012.json"
FORMATS = {"docx": 2, "xlsx": 2, "pptx": 2, "html": 2, "eml": 2}
MAX_FILE_BYTES = 20 * 1024 * 1024
SAMPLE_CELLS = 16


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def verify_corpus(directory: Path, manifest_path: Path = MANIFEST) -> tuple[dict[str, Any], str]:
    """Reject altered, missing, extra, or unsafe paths before any parser import."""
    manifest_bytes = manifest_path.read_bytes()
    manifest = json.loads(manifest_bytes)
    sources = manifest.get("sources")
    if (manifest.get("manifest_version") != 1 or not isinstance(sources, list)
            or len(sources) != 10 or manifest.get("source_file_count") != 10):
        raise ValueError("manifest_shape_mismatch")
    if collections.Counter(item.get("format") for item in sources) != FORMATS:
        raise ValueError("manifest_format_mismatch")
    root = directory.resolve(strict=True)
    if not root.is_dir():
        raise ValueError("corpus_not_directory")
    if str(root).startswith("\\\\"):
        raise ValueError("network_corpus_path_forbidden")
    names: set[str] = set()
    source_ids: set[str] = set()
    for item in sources:
        source_id = item["id"]
        name = item["filename"]
        if source_id in source_ids or name in names:
            raise ValueError("duplicate_source")
        source_ids.add(source_id)
        names.add(name)
        if Path(name).name != name or Path(name).suffix.lower() != "." + item["format"]:
            raise ValueError("unsafe_source_name")
        path = root / name
        if path.is_symlink() or path.resolve(strict=True).parent != root:
            raise ValueError("unsafe_source_path")
        size = path.stat().st_size
        if size != item["bytes"] or size > MAX_FILE_BYTES or _sha256(path) != item["sha256"]:
            raise ValueError(f"source_hash_mismatch:{source_id}")
        if item["format"] == "eml":
            raw_name = item["raw_filename"]
            if Path(raw_name).name != raw_name or raw_name in names:
                raise ValueError("unsafe_raw_source_name")
            names.add(raw_name)
            raw_path = root / raw_name
            if raw_path.is_symlink() or raw_path.resolve(strict=True).parent != root:
                raise ValueError("unsafe_raw_source_path")
            raw = raw_path.read_bytes()
            if (len(raw) != item["raw_bytes"] or len(raw) > MAX_FILE_BYTES
                    or hashlib.sha256(raw).hexdigest() != item["raw_sha256"]):
                raise ValueError(f"raw_source_hash_mismatch:{source_id}")
            if not raw.startswith(b"From ") or b"\n" not in raw or raw[raw.index(b"\n") + 1:] != path.read_bytes():
                raise ValueError(f"eml_derivation_mismatch:{source_id}")
    if {path.name for path in root.iterdir()} != names:
        raise ValueError("unexpected_corpus_file")
    return manifest, hashlib.sha256(manifest_bytes).hexdigest()


def sample_indices(count: int, maximum: int = SAMPLE_CELLS) -> list[int]:
    """Evenly sample from all table cells, including both ends, without cherry-picking."""
    if count < 0 or maximum < 1:
        raise ValueError("invalid_sample_size")
    if count <= maximum:
        return list(range(count))
    if maximum == 1:
        return [0]
    return [(index * (count - 1)) // (maximum - 1) for index in range(maximum)]


def _html_cell(root: Any, locator: dict[str, Any]) -> str:
    matches = root.getroottree().xpath(locator["xpath"])
    if len(matches) != 1 or matches[0].tag not in {"th", "td"}:
        raise ValueError("html_locator_not_unique_cell")

    def hidden(element: Any) -> bool:
        if not isinstance(element.tag, str):
            return True
        tag = str(element.tag).lower()
        style = element.get("style") or ""
        return (tag in {"script", "style", "noscript", "template", "svg", "canvas"}
                or element.get("hidden") is not None
                or (element.get("aria-hidden") or "").lower() == "true"
                or (tag == "input" and (element.get("type") or "").lower() == "hidden")
                or bool(re.search(
                    r"(?:^|;)\s*(?:display\s*:\s*none|visibility\s*:\s*hidden)\s*(?:!important)?\s*(?:;|$)",
                    style, re.IGNORECASE,
                )))

    def text(node: Any) -> str:
        if hidden(node):
            return ""
        parts = [node.text or ""]
        for child in node:
            parts.append(text(child))
            parts.append(child.tail or "")
        return "".join(parts)

    return text(matches[0]).strip()


def resolve_sample(path: Path, selected: list[tuple[int, dict[str, Any], str]]) -> tuple[int, list[dict[str, Any]]]:
    """Reopen the original once with a native reader and resolve selected cell locators."""
    if not selected:
        return 0, []
    suffix = path.suffix.lower()
    close = lambda: None
    if suffix == ".docx":
        from docx import Document as WordDocument
        from docx.table import Table as WordTable

        blocks = list(WordDocument(path).iter_inner_content())

        def resolve(locator: dict[str, Any]) -> str:
            block_index = locator["block_index"]
            block = blocks[block_index]
            if (locator["kind"] != "docx_table_cell" or not isinstance(block, WordTable)
                    or sum(isinstance(part, WordTable) for part in blocks[:block_index]) != locator["table_index"]):
                raise ValueError("docx_locator_not_table")
            return block.cell(locator["row"], locator["column"]).text

    elif suffix == ".xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries

        workbook = load_workbook(path, read_only=False, data_only=False,
                                 keep_links=False, keep_vba=False)
        close = workbook.close

        def resolve(locator: dict[str, Any]) -> str:
            if locator["kind"] != "xlsx_cell":
                raise ValueError("xlsx_locator_kind")
            sheet = workbook[locator["sheet"]]
            cell = sheet[locator["address"]]
            if (cell.data_type == "f") != locator["formula"]:
                raise ValueError("xlsx_formula_flag_mismatch")
            table_name = locator["table"]
            if table_name is not None:
                if table_name not in sheet.tables:
                    raise ValueError("xlsx_declared_table_missing")
                min_col, min_row, max_col, max_row = range_boundaries(sheet.tables[table_name].ref)
                if not (min_row <= cell.row <= max_row and min_col <= cell.column <= max_col):
                    raise ValueError("xlsx_address_outside_table")
            return "" if cell.value is None else str(cell.value)

    elif suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)

        def resolve(locator: dict[str, Any]) -> str:
            if locator["kind"] != "pptx_table_cell":
                raise ValueError("pptx_locator_kind")
            slide = presentation.slides[locator["slide"] - 1]
            shapes = [shape for shape in slide.shapes if shape.shape_id == locator["shape_id"]]
            if len(shapes) != 1 or not shapes[0].has_table:
                raise ValueError("pptx_shape_not_unique_table")
            return shapes[0].table.cell(locator["row"], locator["column"]).text

    elif suffix in {".html", ".eml"}:
        from lxml import etree

        parser = etree.HTMLParser(no_network=True, recover=True, huge_tree=False)
        if suffix == ".html":
            root = etree.fromstring(path.read_bytes(), parser=parser)

            def resolve(locator: dict[str, Any]) -> str:
                if locator["kind"] != "html_dom_cell":
                    raise ValueError("html_locator_kind")
                return _html_cell(root, locator)

        else:
            from email import policy
            from email.parser import BytesParser

            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            roots: dict[tuple[int, ...], Any] = {}

            def resolve(locator: dict[str, Any]) -> str:
                if locator["kind"] != "eml_html_cell":
                    raise ValueError("eml_locator_kind")
                mime_path = tuple(locator["mime_path"])
                if mime_path not in roots:
                    part = message
                    for index in mime_path:
                        part = list(part.iter_parts())[index]
                    if part.get_content_type() != "text/html":
                        raise ValueError("eml_locator_not_html")
                    body = part.get_payload(decode=True)
                    if not isinstance(body, bytes):
                        raise ValueError("eml_invalid_html_body")
                    roots[mime_path] = etree.fromstring(body, parser=parser)
                return _html_cell(roots[mime_path], locator)

    else:
        raise ValueError("unsupported_format")

    matches = 0
    errors = []
    try:
        for index, locator, expected in selected:
            try:
                if resolve(locator) == expected:
                    matches += 1
                else:
                    errors.append({"cell_index": index, "reason": "source_text_mismatch"})
            except (KeyError, IndexError, TypeError, ValueError, AttributeError) as exc:
                errors.append({"cell_index": index, "reason": type(exc).__name__})
    finally:
        close()
    return matches, errors


def inspect_parsed(path: Path, document: Any, parsed: Any, expected_sha256: str) -> dict[str, Any]:
    """Count full-artifact integrity and sample original-source locator resolution."""
    if parsed.source_sha256 != expected_sha256 or parsed.source_format != path.suffix[1:].lower():
        raise ValueError("parsed_source_identity_mismatch")
    units = {unit.unit_id: unit for unit in parsed.units}
    page_text_matches = len(document.pages) == len(parsed.units) and all(
        page.text == unit.text and page.metadata.get("native_unit_id") == unit.unit_id
        for page, unit in zip(document.pages, parsed.units)
    )
    cells = [cell for table in parsed.tables for cell in table.cells]
    locator_count = 0
    span_count = 0
    span_exact = 0
    nonempty_count = 0
    selected = []
    sampled = set(sample_indices(len(cells)))
    for index, cell in enumerate(cells):
        if cell.text:
            nonempty_count += 1
        anchor = cell.anchors[0] if cell.anchors else None
        locator = None
        if anchor is not None:
            if anchor.native_locator:
                try:
                    locator = json.loads(anchor.native_locator)
                    if not isinstance(locator, dict):
                        locator = None
                    else:
                        locator_count += 1
                except json.JSONDecodeError:
                    pass
            if anchor.text_span is not None:
                span_count += 1
                unit = units.get(anchor.unit_id)
                if (unit is not None and unit.text[slice(*anchor.text_span)] == cell.text
                        and anchor.quote == cell.text):
                    span_exact += 1
        if index in sampled and locator is not None:
            selected.append((index, locator, cell.text))
    resolve_started = time.perf_counter()
    sample_exact, sample_errors = resolve_sample(path, selected)
    selected_ids = {index for index, _, _ in selected}
    sample_errors.extend({"cell_index": index, "reason": "missing_native_locator"}
                         for index in sampled - selected_ids)
    sample_errors.sort(key=lambda item: item["cell_index"])
    exact_integrity = (page_text_matches and locator_count == len(cells)
                       and span_exact == span_count == nonempty_count
                       and sample_exact == len(sampled) and not sample_errors)
    return {
        "status": "parsed" if exact_integrity else "parsed_with_integrity_mismatch",
        "source_units": len(parsed.units), "source_text_chars": sum(len(unit.text) for unit in parsed.units),
        "nonempty_source_units": sum(bool(unit.text.strip()) for unit in parsed.units),
        "tables": len(parsed.tables), "cells": len(cells), "nonempty_cells": nonempty_count,
        "cells_with_json_native_locator": locator_count, "cells_with_text_span": span_count,
        "cells_with_exact_text_span": span_exact,
        "cells_with_raw_box": sum(any(a.raw_box is not None for a in cell.anchors) for cell in cells),
        "cells_with_display_box": sum(any(a.display_box is not None for a in cell.anchors) for cell in cells),
        "sampled_cells": len(sampled), "sampled_cells_resolved_exact": sample_exact,
        "sample_errors": sample_errors, "page_text_matches_units": page_text_matches,
        "resolver_seconds": round(time.perf_counter() - resolve_started, 4),
        "parser_id": parsed.parser_id, "parser_version": parsed.parser_version,
    }


def worker(path: Path, expected_sha256: str) -> dict[str, Any]:
    """Run the product native parser in one fresh process, without product extraction."""
    import psutil

    from idpo.product_auto.native_formats import NativeLimitError, parse_native_document

    if _sha256(path) != expected_sha256:
        raise ValueError("source_hash_mismatch")
    process = psutil.Process()
    peak_rss = process.memory_info().rss
    stop = threading.Event()

    def sample_memory() -> None:
        nonlocal peak_rss
        while not stop.wait(0.02):
            peak_rss = max(peak_rss, process.memory_info().rss)

    sampler = threading.Thread(target=sample_memory, daemon=True)
    sampler.start()
    started = time.perf_counter()
    try:
        parse_started = time.perf_counter()
        try:
            document, parsed = parse_native_document(path, expected_sha256)
            parse_seconds = round(time.perf_counter() - parse_started, 4)
            result = inspect_parsed(path, document, parsed, expected_sha256)
            result["parse_seconds"] = parse_seconds
        except NativeLimitError as exc:
            result = {"status": "unsupported", "reason": str(exc)}
        except Exception as exc:  # noqa: BLE001 - seal per-file parser failures without aborting the cohort
            result = {"status": "failed", "reason": type(exc).__name__}
    finally:
        stop.set()
        sampler.join(timeout=1)
    result["worker_seconds"] = round(time.perf_counter() - started, 4)
    result["peak_worker_rss_mib_sampled"] = round(max(peak_rss, process.memory_info().rss) / 1024**2, 2)
    return result


def _run_worker(path: Path, source_hash: str, timeout_seconds: int,
                max_rss_mib: int) -> dict[str, Any]:
    import psutil

    command = [sys.executable, "-m", "product_eval.real_native_parser_diagnostic", "worker",
               "--source", str(path), "--sha256", source_hash]
    started = time.perf_counter()
    process = subprocess.Popen(command, cwd=Path(__file__).resolve().parents[1],
                               stdout=subprocess.PIPE, stderr=subprocess.PIPE,
                               text=True, encoding="utf-8")
    try:
        monitored = psutil.Process(process.pid)
    except psutil.NoSuchProcess:
        monitored = None
    peak_rss = 0
    stop_reason = None
    while process.poll() is None:
        elapsed = time.perf_counter() - started
        if monitored is not None:
            try:
                peak_rss = max(peak_rss, monitored.memory_info().rss)
            except psutil.NoSuchProcess:
                pass
        if elapsed > timeout_seconds:
            stop_reason = "timeout"
            process.kill()
            break
        if peak_rss > max_rss_mib * 1024**2:
            stop_reason = "sampled_rss_limit"
            process.kill()
            break
        time.sleep(0.05)
    stdout, _stderr = process.communicate()
    if stop_reason:
        record = {"status": stop_reason}
    elif process.returncode != 0:
        record = {"status": "worker_failed", "reason": f"exit_{process.returncode}"}
    else:
        try:
            record = json.loads(stdout)
        except json.JSONDecodeError:
            record = {"status": "worker_failed", "reason": "invalid_worker_json"}
    record["process_seconds"] = round(time.perf_counter() - started, 4)
    record["parent_peak_rss_mib_sampled"] = round(peak_rss / 1024**2, 2)
    return record


def evaluate(directory: Path, output: Path, *, manifest_path: Path = MANIFEST,
             include_eml_local_only: bool = False, timeout_seconds: int = 120,
             max_rss_mib: int = 1024, min_available_mib: int = 3072) -> dict[str, Any]:
    if not include_eml_local_only:
        raise ValueError("explicit_local_eml_opt_in_required")
    if timeout_seconds < 1 or max_rss_mib < 128 or min_available_mib < 0:
        raise ValueError("invalid_resource_bound")
    import psutil

    manifest, manifest_sha256 = verify_corpus(directory, manifest_path)
    if psutil.virtual_memory().available < min_available_mib * 1024**2:
        raise ValueError("insufficient_available_memory")
    output = output.resolve()
    if output.is_relative_to(Path(__file__).resolve().parents[1]):
        raise ValueError("report_must_be_outside_repository")
    if str(output).startswith("\\\\"):
        raise ValueError("network_report_path_forbidden")
    records = []
    for item in manifest["sources"]:
        if psutil.virtual_memory().available < min_available_mib * 1024**2:
            records.append({"source_id": item["id"], "format": item["format"],
                            "status": "not_run_low_memory"})
            continue
        record = _run_worker(directory / item["filename"], item["sha256"],
                             timeout_seconds, max_rss_mib)
        record.update({"source_id": item["id"], "format": item["format"],
                       "source_sha256": item["sha256"]})
        records.append(record)
    result = {
        "protocol": "real_native_parser_only_html_recover_cap20k_arm_v1",
        "arm_id": "html_recover_cap20k_20261012", "manifest_sha256": manifest_sha256,
        "diagnostic_sha256": _sha256(Path(__file__)),
        "runtime": {"python": sys.version.split()[0], "platform": platform.platform(),
                    "packages": {name: metadata.version(name) for name in
                                 ("python-docx", "openpyxl", "python-pptx", "lxml", "psutil")}},
        "limits": {"single_worker": True, "timeout_seconds_per_file": timeout_seconds,
                   "sampled_rss_limit_mib": max_rss_mib,
                   "minimum_available_memory_mib": min_available_mib,
                   "maximum_locator_sample_per_file": SAMPLE_CELLS},
        "eml_policy": "local_research_only_no_cloud_no_redistribution",
        "gold_labels_used": False, "model_calls": 0, "network_calls": 0,
        "records": records,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    temporary = output.with_suffix(output.suffix + ".tmp")
    temporary.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    os.replace(temporary, output)
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    parent = commands.add_parser("evaluate")
    parent.add_argument("--corpus-dir", type=Path, required=True)
    parent.add_argument("--output", type=Path, required=True)
    parent.add_argument("--manifest", type=Path, default=MANIFEST)
    parent.add_argument("--include-eml-local-only", action="store_true")
    parent.add_argument("--timeout-seconds", type=int, default=120)
    parent.add_argument("--max-rss-mib", type=int, default=1024)
    parent.add_argument("--min-available-mib", type=int, default=3072)
    child = commands.add_parser("worker")
    child.add_argument("--source", type=Path, required=True)
    child.add_argument("--sha256", required=True)
    args = parser.parse_args()
    if args.command == "evaluate":
        evaluate(args.corpus_dir, args.output, manifest_path=args.manifest,
                 include_eml_local_only=args.include_eml_local_only,
                 timeout_seconds=args.timeout_seconds, max_rss_mib=args.max_rss_mib,
                 min_available_mib=args.min_available_mib)
    else:
        print(json.dumps(worker(args.source, args.sha256), ensure_ascii=False))


if __name__ == "__main__":
    main()
