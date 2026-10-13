"""Research-only, source-verified table graph for the five native-format fixtures.

This module neither reads labels nor calls a model. A candidate is a pointer to
literal parser cells, never an approved field or a claim about business meaning.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import re
from collections.abc import Callable, Iterator
from contextlib import contextmanager
from dataclasses import asdict, dataclass
from email import policy
from email.parser import BytesParser
from pathlib import Path

from idpo.product_auto.native_formats import parse_native_document
from idpo.product_auto.parsed_document import (
    ParsedDocument,
    SourceAnchor,
    Table,
    TableCell,
)

FORMATS = {"docx", "xlsx", "pptx", "html", "eml"}
MAX_SOURCE_BYTES = 20 * 1024 * 1024
MAX_TABLES = 16
MAX_CELLS = 256
MAX_COLUMNS = 12
MAX_ROWS = 64
MAX_CELL_CHARS = 500
MAX_MODEL_INPUT_BYTES = 24_000
FIXTURE_MANIFEST_SHA256 = "dca2d96ecb8e863d4081f113979d0175564968a402eba3b2fe16b40a9743ca23"
FIXTURE_NAMES = ("order.docx", "order.xlsx", "order.html", "order.eml", "order.pptx")


class SourceGraphError(ValueError):
    """A source-wide failure; the code is safe for a source-free receipt."""


class _RejectTable(ValueError):
    pass


def _canonical(value: object) -> str:
    return json.dumps(value, sort_keys=True, separators=(",", ":"), ensure_ascii=False)


def _identity(prefix: str, value: object) -> str:
    return f"{prefix}_{hashlib.sha256(_canonical(value).encode('utf-8')).hexdigest()[:24]}"


def _sha256(path: Path) -> str:
    digest = hashlib.sha256()
    with path.open("rb") as source:
        for chunk in iter(lambda: source.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


@dataclass(frozen=True)
class SourceNode:
    node_id: str
    source_sha256: str
    parser_id: str
    parser_version: str
    parser_config_sha256: str | None
    table_id: str
    unit_id: str
    row: int
    column: int
    text: str
    text_span: tuple[int, int]
    native_locator: str


@dataclass(frozen=True)
class ColumnCandidate:
    candidate_id: str
    header_node_id: str
    label_hint: str
    column: int


@dataclass(frozen=True)
class RowCandidate:
    row_id: str
    source_row: int
    cell_node_ids: tuple[str, ...]


@dataclass(frozen=True)
class TableCandidate:
    table_id: str
    header_basis: str
    columns: tuple[ColumnCandidate, ...]
    rows: tuple[RowCandidate, ...]


@dataclass(frozen=True)
class TableAbstention:
    table_id: str
    reason: str


@dataclass(frozen=True)
class SourceGraph:
    source_sha256: str
    source_format: str
    parser_id: str
    parser_version: str
    parser_config_sha256: str | None
    nodes: tuple[SourceNode, ...]
    tables: tuple[TableCandidate, ...]
    abstentions: tuple[TableAbstention, ...]

    def model_input(self) -> bytes:
        """Serialize atomically. Oversize input abstains instead of truncating."""
        nodes = {node.node_id: node for node in self.nodes}
        if len(nodes) != len(self.nodes):
            raise SourceGraphError("duplicate_node_id")
        if any(node.source_sha256 != self.source_sha256
               or node.parser_id != self.parser_id
               or node.parser_version != self.parser_version
               or node.parser_config_sha256 != self.parser_config_sha256
               for node in self.nodes):
            raise SourceGraphError("node_source_mismatch")
        for table in self.tables:
            for column in table.columns:
                node = nodes.get(column.header_node_id)
                if node is None or node.table_id != table.table_id or node.column != column.column:
                    raise SourceGraphError("invalid_candidate_pointer")
            for row in table.rows:
                if len(row.cell_node_ids) != len(table.columns):
                    raise SourceGraphError("invalid_candidate_pointer")
                for column, node_id in zip(table.columns, row.cell_node_ids):
                    node = nodes.get(node_id)
                    if (node is None or node.table_id != table.table_id
                            or node.row != row.source_row or node.column != column.column):
                        raise SourceGraphError("invalid_candidate_pointer")
        payload = _canonical({
            "source_sha256": self.source_sha256,
            "source_format": self.source_format,
            "parser": {"id": self.parser_id, "version": self.parser_version},
            "parser_config_sha256": self.parser_config_sha256,
            "nodes": [asdict(node) for node in self.nodes],
            "tables": [asdict(table) for table in self.tables],
            "abstentions": [asdict(item) for item in self.abstentions],
        }).encode("utf-8")
        if len(payload) > MAX_MODEL_INPUT_BYTES:
            raise SourceGraphError("model_input_limit")
        return payload


def _native_locator(anchor: SourceAnchor) -> dict[str, object]:
    if not anchor.native_locator:
        raise _RejectTable("missing_native_locator")
    try:
        locator = json.loads(anchor.native_locator)
    except ValueError as exc:
        raise _RejectTable("invalid_native_locator") from exc
    if not isinstance(locator, dict) or _canonical(locator) != anchor.native_locator:
        raise _RejectTable("invalid_native_locator")
    return locator


def _label_hint(text: str) -> str:
    label = text.strip()
    if not label.isascii():
        raise _RejectTable("unsupported_header_script")
    if (not 2 <= len(label) <= 40 or len(label.split()) > 6
            or not re.search(r"[A-Za-z]", label)):
        raise _RejectTable("ambiguous_header")
    hint = re.sub(r"[^a-z0-9]+", "_", label.casefold()).strip("_")
    if not 2 <= len(hint) <= 64:
        raise _RejectTable("ambiguous_header")
    return hint


def _grid(table: Table) -> list[list[TableCell]]:
    cells = table.cells
    if not cells or len(cells) > MAX_CELLS:
        raise _RejectTable("table_cell_limit_or_empty")
    if any(cell.row_span != 1 or cell.column_span != 1 for cell in cells):
        raise _RejectTable("spanned_cell")
    rows = sorted({cell.row for cell in cells})
    columns = sorted({cell.column for cell in cells})
    if not (2 <= len(rows) <= MAX_ROWS and 2 <= len(columns) <= MAX_COLUMNS):
        raise _RejectTable("table_shape_limit")
    if (rows != list(range(rows[0], rows[-1] + 1))
            or columns != list(range(columns[0], columns[-1] + 1))
            or len(cells) != len(rows) * len(columns)):
        raise _RejectTable("incomplete_grid")
    positions = {(cell.row, cell.column): cell for cell in cells}
    if len(positions) != len(cells):
        raise _RejectTable("duplicate_grid_position")
    grid = [[positions[row, column] for column in columns] for row in rows]
    hints = [_label_hint(cell.text) for cell in grid[0]]
    if len(set(hints)) != len(hints):
        raise _RejectTable("duplicate_header")
    if any(not cell.text.strip() or len(cell.text) > MAX_CELL_CHARS
           for row in grid for cell in row):
        raise _RejectTable("blank_or_long_cell")
    if any([cell.text.strip().casefold() for cell in row]
           == [cell.text.strip().casefold() for cell in grid[0]] for row in grid[1:]):
        raise _RejectTable("repeated_header")
    return grid


def _html_content(path: Path, source_format: str):
    from lxml import etree

    source = path.read_bytes()
    selected_path: list[int] | None = None
    encoding: str | None = None
    if source_format == "eml":
        message = BytesParser(policy=policy.default).parsebytes(source)
        bodies = []

        def visit(part, mime_path: list[int]) -> None:
            if part.defects:
                raise SourceGraphError("parser_failed")
            if part.is_multipart():
                for index, child in enumerate(part.iter_parts()):
                    visit(child, [*mime_path, index])
            elif (part.get_content_disposition() != "attachment"
                  and part.get_content_type() == "text/html"):
                bodies.append((mime_path, part))

        visit(message, [])
        if not bodies:
            raise SourceGraphError("html_body_unavailable")
        selected_path, part = bodies[0]
        payload = part.get_payload(decode=True)
        if not isinstance(payload, bytes):
            raise SourceGraphError("parser_failed")
        charset = part.get_content_charset()
        if charset is None:
            source = payload
        else:
            try:
                source = payload.decode(charset, errors="strict").encode("utf-8")
            except (LookupError, UnicodeError) as exc:
                raise SourceGraphError("parser_failed") from exc
            encoding = "utf-8"
    root = etree.fromstring(source, parser=etree.HTMLParser(
        encoding=encoding, no_network=True, recover=True, huge_tree=False))
    if root is None:
        raise SourceGraphError("parser_failed")
    return root.getroottree(), selected_path


@contextmanager
def _resolver(path: Path, source_format: str) -> Iterator[Callable[[TableCell, str], tuple[str, str]]]:
    """Return source text and header basis for an exact native cell locator."""
    if source_format == "docx":
        from docx import Document
        from docx.table import Table as WordTable

        blocks = list(Document(path).iter_inner_content())

        def resolve(cell: TableCell, table_id: str) -> tuple[str, str]:
            locator = _native_locator(cell.anchors[0])
            block_index = locator.get("block_index")
            table_index = locator.get("table_index")
            if (locator.get("kind") != "docx_table_cell"
                    or type(block_index) is not int or not 0 <= block_index < len(blocks)
                    or type(table_index) is not int or table_id != f"table-{table_index + 1}"
                    or sum(isinstance(block, WordTable) for block in blocks[:block_index]) != table_index
                    or not isinstance(blocks[block_index], WordTable)
                    or locator.get("row") != cell.row or locator.get("column") != cell.column):
                raise _RejectTable("native_locator_mismatch")
            return blocks[block_index].cell(cell.row, cell.column).text, "first_row_candidate"

        yield resolve
    elif source_format == "xlsx":
        from openpyxl import load_workbook
        from openpyxl.utils.cell import range_boundaries

        workbook = load_workbook(path, read_only=False, data_only=False,
                                 keep_links=False, keep_vba=False)
        try:
            def resolve(cell: TableCell, table_id: str) -> tuple[str, str]:
                locator = _native_locator(cell.anchors[0])
                sheet_name = locator.get("sheet")
                table_name = locator.get("table")
                if (locator.get("kind") != "xlsx_cell"
                        or not isinstance(sheet_name, str) or sheet_name not in workbook
                        or not isinstance(table_name, str) or not table_name
                        or locator.get("formula") is not False
                        or locator.get("row") is not None
                        or locator.get("address") is None):
                    raise _RejectTable("native_locator_mismatch")
                sheet = workbook[sheet_name]
                sheet_index = workbook.worksheets.index(sheet) + 1
                if (table_id != f"sheet-grid-{sheet_index}" or sheet.sheet_state != "visible"
                        or table_name not in sheet.tables):
                    raise _RejectTable("native_locator_mismatch")
                declared = sheet.tables[table_name]
                if (declared.headerRowCount != 1 or declared.totalsRowCount
                        or declared.totalsRowShown):
                    raise _RejectTable("ambiguous_declared_header")
                first_col, first_row, last_col, last_row = range_boundaries(declared.ref)
                source_cell = sheet.cell(cell.row + 1, cell.column + 1)
                if (not first_row <= source_cell.row <= last_row
                        or not first_col <= source_cell.column <= last_col
                        or locator.get("address") != source_cell.coordinate
                        or source_cell.data_type == "f"):
                    raise _RejectTable("native_locator_mismatch")
                return ("" if source_cell.value is None else str(source_cell.value)), "declared_xlsx_table"

            yield resolve
        finally:
            workbook.close()
    elif source_format == "pptx":
        from pptx import Presentation

        slides = Presentation(path).slides

        def resolve(cell: TableCell, table_id: str) -> tuple[str, str]:
            locator = _native_locator(cell.anchors[0])
            slide_number = locator.get("slide")
            shape_id = locator.get("shape_id")
            if (locator.get("kind") != "pptx_table_cell"
                    or type(slide_number) is not int or not 1 <= slide_number <= len(slides)
                    or type(shape_id) is not int or table_id != f"slide-{slide_number}-shape-{shape_id}"
                    or locator.get("row") != cell.row or locator.get("column") != cell.column):
                raise _RejectTable("native_locator_mismatch")
            shapes = [shape for shape in slides[slide_number - 1].shapes if shape.shape_id == shape_id]
            if len(shapes) != 1 or not shapes[0].has_table:
                raise _RejectTable("native_locator_mismatch")
            return shapes[0].table.cell(cell.row, cell.column).text, "first_row_candidate"

        yield resolve
    else:
        from idpo.product_auto.native_formats import (
            _html_hidden,
            _html_visible_cell_text,
        )

        tree, selected_path = _html_content(path, source_format)
        tables = [table for table in tree.xpath("//body//table")
                  if not any(_html_hidden(item) for item in (table, *table.iterancestors()))]

        def resolve(cell: TableCell, table_id: str) -> tuple[str, str]:
            locator = _native_locator(cell.anchors[0])
            expected_kind = "eml_html_cell" if source_format == "eml" else "html_dom_cell"
            xpath = locator.get("xpath")
            if (locator.get("kind") != expected_kind or not isinstance(xpath, str)
                    or not xpath.startswith("/") or (source_format == "eml"
                    and locator.get("mime_path") != selected_path)):
                raise _RejectTable("native_locator_mismatch")
            matches = tree.xpath(xpath)
            if len(matches) != 1 or matches[0].tag not in {"th", "td"}:
                raise _RejectTable("native_locator_mismatch")
            element = matches[0]
            if any(_html_hidden(item) for item in (element, *element.iterancestors())):
                raise _RejectTable("hidden_cell")
            table = next(iter(element.iterancestors("table")), None)
            if table is None or table not in tables or table_id != f"table-{tables.index(table) + 1}":
                raise _RejectTable("native_locator_mismatch")
            rows = table.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
            rows = [row for row in rows if not any(_html_hidden(item)
                    for item in (row, *row.iterancestors()))]
            if cell.row >= len(rows) or element.getparent() is not rows[cell.row]:
                raise _RejectTable("native_locator_mismatch")
            columns = [item for item in rows[cell.row].xpath("./th | ./td")
                       if not _html_hidden(item)]
            if cell.column >= len(columns) or columns[cell.column] is not element:
                raise _RejectTable("native_locator_mismatch")
            header_basis = ("html_thead" if rows[0].getparent().tag == "thead"
                            and all(item.tag == "th" for item in rows[0].xpath("./th | ./td"))
                            else "first_row_candidate")
            return _html_visible_cell_text(element), header_basis

        yield resolve


def _validated_node(parsed: ParsedDocument, table_id: str, cell: TableCell,
                    resolve: Callable[[TableCell, str], tuple[str, str]]) -> tuple[SourceNode, str]:
    if len(cell.anchors) != 1:
        raise _RejectTable("ambiguous_cell_anchor")
    anchor = cell.anchors[0]
    units = {unit.unit_id: unit for unit in parsed.units}
    unit = units.get(anchor.unit_id)
    if (unit is None or anchor.text_span is None or anchor.quote != cell.text
            or not cell.text or len(cell.text) > MAX_CELL_CHARS):
        raise _RejectTable("invalid_cell_span")
    start, end = anchor.text_span
    if unit.text[start:end] != cell.text:
        raise _RejectTable("invalid_cell_span")
    locator = _native_locator(anchor)
    try:
        source_text, header_basis = resolve(cell, table_id)
    except _RejectTable:
        raise
    except Exception as exc:  # Invalid source location stays an abstention.
        raise _RejectTable("native_locator_unresolvable") from exc
    if source_text != cell.text:
        raise _RejectTable("native_text_mismatch")
    identity = [parsed.source_sha256, parsed.parser_id, parsed.parser_version,
                parsed.parser_config_sha256,
                table_id, anchor.unit_id, cell.row, cell.column,
                anchor.native_locator, anchor.text_span]
    return SourceNode(_identity("n", identity), parsed.source_sha256,
                      parsed.parser_id, parsed.parser_version, parsed.parser_config_sha256,
                      table_id, anchor.unit_id,
                      cell.row, cell.column, cell.text, anchor.text_span,
                      anchor.native_locator or _canonical(locator)), header_basis


def build_source_graph(parsed: ParsedDocument, source_path: Path) -> SourceGraph:
    """Use only parser output and the hashed original; isolate bad tables."""
    if (parsed.source_format not in FORMATS
            or source_path.suffix.lower() != f".{parsed.source_format}"):
        raise SourceGraphError("unsupported_format")
    try:
        if source_path.stat().st_size > MAX_SOURCE_BYTES:
            raise SourceGraphError("source_size_limit")
        if _sha256(source_path) != parsed.source_sha256:
            raise SourceGraphError("source_hash_mismatch")
    except OSError as exc:
        raise SourceGraphError("source_unavailable") from exc
    if len(parsed.tables) > MAX_TABLES or sum(len(table.cells) for table in parsed.tables) > MAX_CELLS:
        raise SourceGraphError("graph_limit")

    nodes: list[SourceNode] = []
    tables: list[TableCandidate] = []
    abstentions: list[TableAbstention] = []
    try:
        with _resolver(source_path, parsed.source_format) as resolve:
            for table in parsed.tables:
                try:
                    grid = _grid(table)
                    table_nodes = [_validated_node(parsed, table.table_id, cell, resolve)
                                   for row in grid for cell in row]
                    if parsed.source_format == "xlsx":
                        from openpyxl.utils.cell import range_boundaries

                        locators = [_native_locator(cell.anchors[0]) for row in grid for cell in row]
                        table_names = {locator.get("table") for locator in locators}
                        sheet_names = {locator.get("sheet") for locator in locators}
                        if len(table_names) != 1 or len(sheet_names) != 1:
                            raise _RejectTable("mixed_declared_table")
                        # The parser's grid may include non-table cells; one complete
                        # declared table is the only safe header/row witness here.
                        from openpyxl import load_workbook

                        book = load_workbook(source_path, read_only=False, data_only=False,
                                             keep_links=False, keep_vba=False)
                        try:
                            declared = book[next(iter(sheet_names))].tables[next(iter(table_names))]
                            c0, r0, c1, r1 = range_boundaries(declared.ref)
                            expected = {(r - 1, c - 1) for r in range(r0, r1 + 1)
                                        for c in range(c0, c1 + 1)}
                            if (expected != {(cell.row, cell.column) for row in grid for cell in row}
                                    or [column.name for column in declared.tableColumns]
                                    != [cell.text for cell in grid[0]]):
                                raise _RejectTable("declared_table_mismatch")
                            sheet = book[next(iter(sheet_names))]
                            if any(sheet.row_dimensions[row].hidden for row in range(r0, r1 + 1)):
                                raise _RejectTable("hidden_declared_row")
                            if any(dimension.hidden and dimension.min and dimension.max
                                   and any(dimension.min <= column <= dimension.max
                                           for column in range(c0, c1 + 1))
                                   for dimension in sheet.column_dimensions.values()):
                                raise _RejectTable("hidden_declared_column")
                        finally:
                            book.close()
                    if len({node.native_locator for node, _ in table_nodes}) != len(table_nodes):
                        raise _RejectTable("duplicate_native_locator")
                    if len({basis for _, basis in table_nodes}) != 1:
                        raise _RejectTable("mixed_header_basis")
                    basis = table_nodes[0][1]
                    matrix = [[table_nodes[index * len(grid[0]) + column][0]
                               for column in range(len(grid[0]))] for index in range(len(grid))]
                    identity = [parsed.source_sha256, parsed.parser_id,
                                parsed.parser_version, parsed.parser_config_sha256,
                                table.table_id]
                    columns = tuple(ColumnCandidate(_identity("c", [*identity, cell.column]),
                                                    cell.node_id, _label_hint(cell.text), cell.column)
                                    for cell in matrix[0])
                    rows = tuple(RowCandidate(_identity("r", [*identity, row[0].row]),
                                              row[0].row, tuple(cell.node_id for cell in row))
                                 for row in matrix[1:])
                    tables.append(TableCandidate(table.table_id, basis, columns, rows))
                    nodes.extend(node for node, _ in table_nodes)
                except _RejectTable as exc:
                    abstentions.append(TableAbstention(table.table_id, str(exc)))
                except Exception:  # noqa: BLE001 - untrusted table fails closed without leaking source text
                    abstentions.append(TableAbstention(table.table_id, "source_validation_failed"))
    except SourceGraphError:
        raise
    except Exception as exc:
        raise SourceGraphError("source_validation_failed") from exc
    if (len({node.node_id for node in nodes}) != len(nodes)
            or len({node.native_locator for node in nodes}) != len(nodes)):
        raise SourceGraphError("duplicate_node_id")
    try:
        if _sha256(source_path) != parsed.source_sha256:
            raise SourceGraphError("source_changed_during_validation")
    except OSError as exc:
        raise SourceGraphError("source_unavailable") from exc
    graph = SourceGraph(parsed.source_sha256, parsed.source_format, parsed.parser_id,
                        parsed.parser_version, parsed.parser_config_sha256,
                        tuple(nodes), tuple(tables), tuple(abstentions))
    graph.model_input()
    return graph


def parse_source_graph(path: Path) -> SourceGraph:
    """Bounded parsing entry point; parser exceptions do not leak source text."""
    if path.suffix.lower().lstrip(".") not in FORMATS:
        raise SourceGraphError("unsupported_format")
    try:
        if path.stat().st_size > MAX_SOURCE_BYTES:
            raise SourceGraphError("source_size_limit")
        source_sha256 = _sha256(path)
    except OSError as exc:
        raise SourceGraphError("source_unavailable") from exc
    try:
        _, parsed = parse_native_document(path, source_sha256)
    except Exception as exc:
        raise SourceGraphError("parser_failed") from exc
    return build_source_graph(parsed, path)


def locked_fixture_receipt(directory: Path) -> dict[str, object]:
    """Only the five self-created source files; never reads a gold annotation."""
    manifest_path = directory / "manifest.json"
    if _sha256(manifest_path) != FIXTURE_MANIFEST_SHA256:
        raise SourceGraphError("fixture_manifest_mismatch")
    files = json.loads(manifest_path.read_text(encoding="utf-8"))["files"]
    if set(files) != set(FIXTURE_NAMES):
        raise SourceGraphError("fixture_set_mismatch")
    records = []
    for name in FIXTURE_NAMES:
        path = directory / name
        if _sha256(path) != files[name]:
            raise SourceGraphError("fixture_source_mismatch")
        try:
            graph = parse_source_graph(path)
            model_input = graph.model_input()
            records.append({
                "name": name,
                "source_sha256": graph.source_sha256,
                "status": "candidate" if graph.tables else "abstained",
                "parser_id": graph.parser_id,
                "parser_version": graph.parser_version,
                "source_nodes": len(graph.nodes),
                "table_candidates": len(graph.tables),
                "column_candidates": sum(len(table.columns) for table in graph.tables),
                "row_candidates": sum(len(table.rows) for table in graph.tables),
                "abstentions": [asdict(item) for item in graph.abstentions],
                "model_input_bytes": len(model_input),
                "model_input_sha256": hashlib.sha256(model_input).hexdigest(),
            })
        except SourceGraphError as exc:
            records.append({"name": name, "source_sha256": files[name],
                            "status": "abstained", "reason": str(exc)})
    return {"protocol": "native_s1_source_graph_v1",
            "fixture_manifest_sha256": FIXTURE_MANIFEST_SHA256,
            "model_calls": 0, "gold_access": False, "product_extraction_calls": 0,
            "scope": "five self-created native files sharing one synthetic table",
            "limitations": [
                "Counts measure source-pointer enumeration, not schema or extraction accuracy.",
                "First-row labels in DOCX and PPTX are provisional; no semantic approval occurs.",
                "No PDF/image OCR, unseen language, long document, real family, or HITL cost tested.",
            ],
            "records": records}


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("fixture_directory", type=Path)
    args = parser.parse_args()
    print(json.dumps(locked_fixture_receipt(args.fixture_directory),
                     indent=2, ensure_ascii=False))


if __name__ == "__main__":
    main()
