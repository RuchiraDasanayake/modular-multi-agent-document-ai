"""Source-locked native-locator diagnostic for the five synthetic office/web/email fixtures.

This is parser feasibility only. It does not run product extraction or an AI model.
"""

from __future__ import annotations

import argparse
import hashlib
import json
import platform
import subprocess
import sys
import threading
import time
from email import policy
from email.message import EmailMessage
from email.parser import BytesParser
from importlib import metadata
from pathlib import Path
from typing import Any

from product_eval.mixed_format_docling_diagnostic import (
    FIXTURE_NAMES,
    HEADERS,
    ROWS,
    verify_fixtures,
)


MAX_SOURCE_BYTES = 20 * 1024 * 1024
EXPECTED_GRID = (HEADERS, *ROWS)


def _cell(text: str, locator: dict[str, Any]) -> dict[str, Any]:
    return {"text": text, "locator": locator, "cell_bbox": None}


def _html_cells(source: bytes, *, mime_path: list[int] | None = None) -> list[dict[str, Any]]:
    from lxml import etree

    parser = etree.HTMLParser(no_network=True, recover=False, huge_tree=False)
    root = etree.fromstring(source, parser=parser)
    tree = root.getroottree()
    cells = []
    for table in root.xpath("//table"):
        # This direct-child traversal avoids counting a nested table's rows twice.
        rows = table.xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
        for row in rows:
            for item in row.xpath("./th | ./td"):
                locator: dict[str, Any] = {
                    "kind": "html_dom_cell", "xpath": tree.getpath(item),
                }
                if mime_path is not None:
                    locator = {"kind": "eml_html_cell", "mime_path": mime_path,
                               "xpath": tree.getpath(item)}
                cells.append(_cell("".join(item.itertext()).strip(), locator))
    return cells


def _mime_part_path(message: EmailMessage, selected: EmailMessage) -> list[int]:
    def visit(current: EmailMessage, path: list[int]) -> list[int] | None:
        if current is selected:
            return path
        if current.is_multipart():
            for index, child in enumerate(current.iter_parts()):
                found = visit(child, [*path, index])
                if found is not None:
                    return found
        return None

    result = visit(message, [])
    if result is None:
        raise ValueError("Selected HTML body is not in the MIME tree")
    return result


def parse_source(path: Path) -> dict[str, Any]:
    suffix = path.suffix.lower()
    if suffix == ".docx":
        from docx import Document
        from docx.table import Table

        document = Document(path)
        cells = []
        paragraphs = []
        table_index = 0
        for block_index, block in enumerate(document.iter_inner_content()):
            if isinstance(block, Table):
                for row_index, row in enumerate(block.rows):
                    for column_index, table_cell in enumerate(row.cells):
                        cells.append(_cell(table_cell.text, {
                            "kind": "docx_table_cell", "block_index": block_index,
                            "table_index": table_index, "row": row_index,
                            "column": column_index,
                        }))
                table_index += 1
            else:
                paragraphs.append({"text": block.text,
                                   "locator": {"kind": "docx_paragraph",
                                               "block_index": block_index}})
        return {"cells": cells, "paragraphs": paragraphs,
                "tables": table_index, "source_native_geometry": "none"}

    if suffix == ".xlsx":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=False, data_only=False,
                                 keep_links=False, keep_vba=False)
        cells = []
        tables = []
        try:
            for sheet in workbook.worksheets:
                for table in sheet.tables.values():
                    tables.append({"sheet": sheet.title, "name": table.name, "ref": table.ref})
                    for row in sheet[table.ref]:
                        for sheet_cell in row:
                            value = sheet_cell.value
                            cells.append(_cell("" if value is None else str(value), {
                                "kind": "xlsx_cell", "sheet": sheet.title,
                                "table": table.name, "address": sheet_cell.coordinate,
                                "formula": sheet_cell.data_type == "f",
                            }))
        finally:
            workbook.close()
        return {"cells": cells, "tables": tables,
                "source_native_geometry": "sheet_cell_address_only"}

    if suffix == ".pptx":
        from pptx import Presentation

        presentation = Presentation(path)
        cells = []
        tables = []
        for slide_index, slide in enumerate(presentation.slides):
            for shape in slide.shapes:
                if not shape.has_table:
                    continue
                table = shape.table
                tables.append({"slide": slide_index + 1, "shape_id": shape.shape_id,
                               "table_region_emu": {"left": shape.left, "top": shape.top,
                                                    "width": shape.width,
                                                    "height": shape.height}})
                for row_index in range(len(table.rows)):
                    for column_index in range(len(table.columns)):
                        ppt_cell = table.cell(row_index, column_index)
                        cells.append(_cell(ppt_cell.text, {
                            "kind": "pptx_table_cell", "slide": slide_index + 1,
                            "shape_id": shape.shape_id, "row": row_index,
                            "column": column_index,
                            "is_spanned": ppt_cell.is_spanned,
                        }))
        return {"cells": cells, "tables": tables,
                "source_native_geometry": "table_region_emu_only"}

    if suffix == ".html":
        return {"cells": _html_cells(path.read_bytes()),
                "source_native_geometry": "dom_xpath_only"}

    if suffix == ".eml":
        message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
        if message.defects:
            raise ValueError(f"Malformed email: {message.defects!r}")
        selected = message.get_body(preferencelist=("html", "plain"))
        if selected is None or selected.get_content_type() != "text/html":
            return {"cells": [], "selected_body_type": selected.get_content_type()
                    if selected is not None else None,
                    "source_native_geometry": "mime_path_only"}
        mime_path = _mime_part_path(message, selected)
        content = selected.get_content()
        return {"cells": _html_cells(content.encode("utf-8"), mime_path=mime_path),
                "selected_body_type": "text/html", "selected_mime_path": mime_path,
                "subject": message.get("Subject"),
                "source_native_geometry": "mime_path_plus_dom_xpath_only"}

    raise ValueError(f"Unsupported fixture suffix: {suffix}")


def resolve_locator(path: Path, locator: dict[str, Any]) -> str:
    kind = locator["kind"]
    if kind in {"html_dom_cell", "eml_html_cell"}:
        from lxml import etree

        if kind == "eml_html_cell":
            message = BytesParser(policy=policy.default).parsebytes(path.read_bytes())
            part = message
            for index in locator["mime_path"]:
                part = list(part.iter_parts())[index]
            if part.get_content_type() != "text/html":
                raise ValueError("MIME locator no longer points to HTML")
            source = part.get_content().encode("utf-8")
        else:
            source = path.read_bytes()
        root = etree.fromstring(source, parser=etree.HTMLParser(
            no_network=True, recover=False, huge_tree=False))
        matches = root.getroottree().xpath(locator["xpath"])
        if len(matches) != 1 or matches[0].tag not in {"th", "td"}:
            raise ValueError("DOM locator is not a unique table cell")
        return "".join(matches[0].itertext()).strip()

    if kind in {"docx_table_cell", "docx_paragraph"}:
        from docx import Document
        from docx.table import Table

        blocks = list(Document(path).iter_inner_content())
        block = blocks[locator["block_index"]]
        if kind == "docx_paragraph":
            if isinstance(block, Table):
                raise ValueError("DOCX paragraph locator points to a table")
            return block.text
        if not isinstance(block, Table):
            raise ValueError("DOCX table locator points to a paragraph")
        return block.cell(locator["row"], locator["column"]).text

    if kind == "xlsx_cell":
        from openpyxl import load_workbook

        workbook = load_workbook(path, read_only=False, data_only=False,
                                 keep_links=False, keep_vba=False)
        try:
            sheet = workbook[locator["sheet"]]
            if locator["table"] not in sheet.tables:
                raise ValueError("XLSX table no longer exists")
            from openpyxl.utils.cell import range_boundaries

            min_col, min_row, max_col, max_row = range_boundaries(
                sheet.tables[locator["table"]].ref)
            cell = sheet[locator["address"]]
            if not min_row <= cell.row <= max_row or not min_col <= cell.column <= max_col:
                raise ValueError("XLSX address is outside recorded table")
            return "" if cell.value is None else str(cell.value)
        finally:
            workbook.close()

    if kind == "pptx_table_cell":
        from pptx import Presentation

        presentation = Presentation(path)
        slide = presentation.slides[locator["slide"] - 1]
        matches = [shape for shape in slide.shapes if shape.shape_id == locator["shape_id"]]
        if len(matches) != 1 or not matches[0].has_table:
            raise ValueError("PPTX locator does not point to a unique table")
        return matches[0].table.cell(locator["row"], locator["column"]).text

    raise ValueError(f"Unknown locator kind: {kind}")


def _sha256(path: Path) -> str:
    return hashlib.sha256(path.read_bytes()).hexdigest()


def worker(directory: Path, name: str) -> dict[str, Any]:
    import psutil

    if name not in FIXTURE_NAMES:
        raise ValueError("Unknown fixture")
    manifest = verify_fixtures(directory)
    path = directory / name
    if path.stat().st_size > MAX_SOURCE_BYTES:
        raise ValueError("Fixture exceeds diagnostic size cap")

    process = psutil.Process()
    peak_rss = process.memory_info().rss
    stop = threading.Event()

    def sample_memory() -> None:
        nonlocal peak_rss
        while not stop.wait(0.01):
            peak_rss = max(peak_rss, process.memory_info().rss)

    sampler = threading.Thread(target=sample_memory, daemon=True)
    sampler.start()
    started = time.perf_counter()
    try:
        parse_started = time.perf_counter()
        parsed = parse_source(path)
        parsed["parse_seconds"] = round(time.perf_counter() - parse_started, 4)
        cells = parsed["cells"]
        expected = [item for row in EXPECTED_GRID for item in row]
        exact = sum(cell["text"] == value for cell, value in zip(cells, expected))
        resolved = sum(resolve_locator(path, cell["locator"]) == cell["text"]
                       for cell in cells)
        fixture_match = (len(cells) == len(expected) and exact == len(expected)
                         and resolved == len(expected))
        parsed.update({"status": "match" if fixture_match else "mismatch",
                       "fixture_exact_match": fixture_match, "source": name,
                       "source_sha256": manifest["files"][name],
                       "expected_cells": len(expected), "cells_found": len(cells),
                       "exact_known_cells": exact, "resolvable_cell_locators": resolved,
                       "cell_bbox_count": sum(cell["cell_bbox"] is not None for cell in cells)})
    finally:
        stop.set()
        sampler.join(timeout=1)
    parsed["worker_wall_seconds"] = round(time.perf_counter() - started, 4)
    parsed["peak_rss_mib"] = round(max(peak_rss, process.memory_info().rss) / 1024**2, 2)
    return parsed


def evaluate(directory: Path, output: Path) -> dict[str, Any]:
    manifest = verify_fixtures(directory)
    records = []
    for name in FIXTURE_NAMES:
        started = time.perf_counter()
        run = subprocess.run([sys.executable, "-m", "product_eval.mixed_format_native_diagnostic",
                              "worker", str(directory.resolve()), name],
                             cwd=Path(__file__).resolve().parents[1],
                             capture_output=True, text=True, timeout=120, check=False)
        if run.returncode:
            raise RuntimeError(f"Native worker failed for {name}: {run.stderr[-1000:]}")
        record = json.loads(run.stdout)
        record["process_wall_seconds"] = round(time.perf_counter() - started, 4)
        records.append(record)
    result = {
        "protocol": "mixed_format_native_v1", "fixture_manifest": manifest,
        "runtime": {
            "python": sys.version.split()[0], "platform": platform.platform(),
            "packages": {name: metadata.version(name) for name in
                         ("python-docx", "openpyxl", "python-pptx", "lxml", "psutil")},
        },
        "gpu_used": False, "model_downloads": False,
        "records": records,
    }
    output.parent.mkdir(parents=True, exist_ok=True)
    output.write_text(json.dumps(result, indent=2, ensure_ascii=False) + "\n", encoding="utf-8")
    return result


def main() -> None:
    parser = argparse.ArgumentParser(description=__doc__)
    commands = parser.add_subparsers(dest="command", required=True)
    run = commands.add_parser("evaluate")
    run.add_argument("directory", type=Path)
    run.add_argument("output", type=Path)
    child = commands.add_parser("worker")
    child.add_argument("directory", type=Path)
    child.add_argument("name", choices=FIXTURE_NAMES)
    args = parser.parse_args()
    if args.command == "evaluate":
        evaluate(args.directory, args.output)
    else:
        print(json.dumps(worker(args.directory, args.name), ensure_ascii=False))


if __name__ == "__main__":
    main()
