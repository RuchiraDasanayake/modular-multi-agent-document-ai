"""Conservative, review-only linked rows from source-native table cells.

This adapter proposes literal cells from a complete parser grid. It never
interprets the business meaning of a header or approves a row relation.
"""

from __future__ import annotations

import hashlib
import json
import re
from collections import Counter, defaultdict
from email import policy
from email.parser import BytesParser
from io import BytesIO
from pathlib import Path
from typing import Any

from idpo.product_auto.logic import (
    MAX_LINKED_ROWS_PER_FIELD,
    infer_type,
    normalize_field_id,
)
from idpo.product_auto.native_formats import _html_visible_cell_text
from idpo.product_auto.parsed_document import ParsedDocument, TableCell

MAX_SOURCE_BYTES = 20 * 1024 * 1024
MAX_TABLE_SEGMENTS = 32
MAX_COLUMNS = 8
MAX_CELL_CHARS = 500
_FOOTER_RE = re.compile(
    r"^(?:sub\s*total|grand\s*total|total|balance|summary)\b", re.IGNORECASE
)
_HIDDEN_STYLE_RE = re.compile(
    r"(?:^|;)\s*(?:display\s*:\s*none|visibility\s*:\s*hidden)\b", re.IGNORECASE
)


class _Abstain(ValueError):
    pass


def _locator(cell: TableCell) -> dict[str, Any]:
    if len(cell.anchors) != 1 or not cell.anchors[0].native_locator:
        raise _Abstain("cell_anchor_ambiguous")
    try:
        value = json.loads(cell.anchors[0].native_locator)
    except ValueError as exc:
        raise _Abstain("cell_locator_invalid") from exc
    if not isinstance(value, dict):
        raise _Abstain("cell_locator_invalid")
    return value


def _shape(value: str) -> str:
    inferred = infer_type(value)
    return "number" if inferred in {"integer", "decimal"} else inferred


def _html_root(source: bytes, *, encoding: str | None = None):
    from lxml import etree

    return etree.fromstring(
        source,
        parser=etree.HTMLParser(
            encoding=encoding,
            no_network=True,
            recover=True,
            huge_tree=False,
        ),
    )


def _selected_email_html(source: bytes) -> tuple[list[int], bytes, str | None]:
    message = BytesParser(policy=policy.default).parsebytes(source)
    bodies: list[tuple[list[int], Any]] = []

    def visit(part, path: list[int]) -> None:
        if part.defects:
            raise _Abstain("mime_defect")
        if part.is_multipart():
            for index, child in enumerate(part.iter_parts()):
                visit(child, [*path, index])
        elif (
            part.get_content_disposition() != "attachment"
            and part.get_content_type() == "text/html"
        ):
            bodies.append((path, part))

    visit(message, [])
    if not bodies:
        raise _Abstain("html_body_unavailable")
    path, part = bodies[0]
    payload = part.get_payload(decode=True)
    if not isinstance(payload, bytes):
        raise _Abstain("html_body_unavailable")
    mime_charset = part.get_content_charset()
    if mime_charset is None:
        return path, payload, None
    try:
        content = payload.decode(mime_charset, errors="strict")
    except (LookupError, UnicodeError) as exc:
        raise _Abstain("invalid_mime_charset") from exc
    return path, content.encode("utf-8"), "utf-8"


def _hidden_html(element) -> bool:
    for ancestor in (element, *element.iterancestors()):
        if str(ancestor.tag).lower() in {"nav", "aside"}:
            return True
        if (
            ancestor.get("hidden") is not None
            or ancestor.get("aria-hidden", "").lower() == "true"
        ):
            return True
        if ancestor.get("role", "").lower() in {"presentation", "none", "navigation"}:
            return True
        if _HIDDEN_STYLE_RE.search(ancestor.get("style", "")):
            return True
    return False


def _prepare_html(source: bytes, source_format: str):
    mime_path: list[int] | None = None
    if source_format == "eml":
        mime_path, html_bytes, encoding = _selected_email_html(source)
        root = _html_root(html_bytes, encoding=encoding)
    else:
        root = _html_root(source)
    return root.getroottree(), mime_path


def _validate_html(cells: list[TableCell], tree, mime_path: list[int] | None) -> str:
    table_path: str | None = None
    for cell in cells:
        locator = _locator(cell)
        expected_kind = "eml_html_cell" if mime_path is not None else "html_dom_cell"
        xpath = locator.get("xpath")
        if (
            locator.get("kind") != expected_kind
            or not isinstance(xpath, str)
            or not xpath.startswith("/")
        ):
            raise _Abstain("html_locator_invalid")
        if mime_path is not None and locator.get("mime_path") != mime_path:
            raise _Abstain("mime_part_mismatch")
        matches = tree.xpath(xpath)
        if len(matches) != 1 or matches[0].tag not in {"th", "td"}:
            raise _Abstain("html_locator_ambiguous")
        element = matches[0]
        source_text = _html_visible_cell_text(element)
        if source_text != cell.text or _hidden_html(element):
            raise _Abstain("html_cell_untrusted")
        ancestors = element.xpath("ancestor::table[1]")
        if len(ancestors) != 1:
            raise _Abstain("html_table_missing")
        current_path = tree.getpath(ancestors[0])
        if table_path is None:
            table_path = current_path
        elif table_path != current_path:
            raise _Abstain("html_table_mixed")
        parent = element.getparent()
        source_rows = [
            row for row in ancestors[0].xpath("./tr | ./thead/tr | ./tbody/tr | ./tfoot/tr")
            if not _hidden_html(row)
        ]
        if (
            parent is None
            or cell.row >= len(source_rows)
            or tree.getpath(source_rows[cell.row]) != tree.getpath(parent)
        ):
            raise _Abstain("html_grid_row_mismatch")
        source_cells = [item for item in parent.xpath("./th | ./td") if not _hidden_html(item)]
        if cell.column >= len(source_cells) or tree.getpath(
            source_cells[cell.column]
        ) != tree.getpath(element):
            raise _Abstain("html_grid_column_mismatch")
        section = parent.getparent() if parent is not None else None
        if cell.row == min(item.row for item in cells):
            if (
                element.tag != "th"
                or parent is None
                or parent.tag != "tr"
                or section is None
                or section.tag != "thead"
            ):
                raise _Abstain("semantic_header_missing")
        elif (
            element.tag != "td"
            or parent is None
            or parent.tag != "tr"
            or section is not None
            and section.tag == "tfoot"
        ):
            raise _Abstain("html_body_cell_invalid")
    return "html_thead"


def _validate_xlsx(cells: list[TableCell], workbook, declared_name: str) -> str:
    from openpyxl.utils.cell import range_boundaries

    locators = [_locator(cell) for cell in cells]
    sheet_names = {locator.get("sheet") for locator in locators}
    if len(sheet_names) != 1 or any(
        locator.get("kind") != "xlsx_cell" or locator.get("table") != declared_name
        for locator in locators
    ):
        raise _Abstain("xlsx_table_locator_mismatch")
    sheet_name = next(iter(sheet_names))
    if not isinstance(sheet_name, str) or sheet_name not in workbook:
        raise _Abstain("xlsx_sheet_missing")
    sheet = workbook[sheet_name]
    if sheet.sheet_state != "visible" or declared_name not in sheet.tables:
        raise _Abstain("xlsx_table_not_visible")
    declared = sheet.tables[declared_name]
    if (
        declared.headerRowCount != 1
        or declared.totalsRowCount
        or declared.totalsRowShown
    ):
        raise _Abstain("xlsx_header_or_total_ambiguous")
    first_col, first_row, last_col, last_row = range_boundaries(declared.ref)
    if {(cell.row, cell.column) for cell in cells} != {
        (row - 1, column - 1)
        for row in range(first_row, last_row + 1)
        for column in range(first_col, last_col + 1)
    }:
        raise _Abstain("xlsx_declared_range_mismatch")
    if any(sheet.row_dimensions[row].hidden for row in range(first_row, last_row + 1)):
        raise _Abstain("xlsx_hidden_row")
    if any(
        dimension.hidden
        and any(
            dimension.min <= column <= dimension.max
            for column in range(first_col, last_col + 1)
        )
        for dimension in sheet.column_dimensions.values()
        if dimension.min and dimension.max
    ):
        raise _Abstain("xlsx_hidden_column")
    headers = [
        cell.text
        for cell in sorted(cells, key=lambda item: item.column)
        if cell.row == first_row - 1
    ]
    if [column.name for column in declared.tableColumns] != headers:
        raise _Abstain("xlsx_header_mismatch")
    for cell, locator in zip(cells, locators):
        source_cell = sheet.cell(cell.row + 1, cell.column + 1)
        if (
            locator.get("address") != source_cell.coordinate
            or locator.get("formula") is not False
            or (source_cell.data_type == "f")
            or ("" if source_cell.value is None else str(source_cell.value))
            != cell.text
        ):
            raise _Abstain("xlsx_cell_mismatch_or_formula")
    return "declared_xlsx_table"


def _validate_office_grid_locators(cells: list[TableCell], source_format: str) -> None:
    expected_kind = "docx_table_cell" if source_format == "docx" else "pptx_table_cell"
    for cell in cells:
        locator = _locator(cell)
        if (
            locator.get("kind") != expected_kind
            or locator.get("row") != cell.row
            or (locator.get("column") != cell.column)
        ):
            raise _Abstain("office_grid_locator_mismatch")
        if source_format == "docx":
            if any(
                type(locator.get(key)) is not int or locator[key] < 0
                for key in ("block_index", "table_index")
            ):
                raise _Abstain("office_grid_locator_mismatch")
        elif (
            type(locator.get("slide")) is not int
            or locator["slide"] < 1
            or (type(locator.get("shape_id")) is not int or locator["shape_id"] < 1)
        ):
            raise _Abstain("office_grid_locator_mismatch")


def _rectangular(
    cells: list[TableCell], source_format: str
) -> tuple[list[list[TableCell]], list[str]]:
    if any(cell.row_span != 1 or cell.column_span != 1 for cell in cells):
        raise _Abstain("spanned_cell")
    rows = sorted({cell.row for cell in cells})
    columns = sorted({cell.column for cell in cells})
    minimum_columns = 3 if source_format in {"docx", "pptx"} else 2
    if not (
        minimum_columns <= len(columns) <= MAX_COLUMNS
        and 3 <= len(rows) <= MAX_LINKED_ROWS_PER_FIELD + 1
    ):
        raise _Abstain("table_shape_outside_limit")
    if rows != list(range(rows[0], rows[-1] + 1)) or columns != list(
        range(columns[0], columns[-1] + 1)
    ):
        raise _Abstain("noncontiguous_grid")
    if len(cells) != len(rows) * len(columns):
        raise _Abstain("incomplete_grid")
    positions = {(cell.row, cell.column): cell for cell in cells}
    if len(positions) != len(cells):
        raise _Abstain("duplicate_grid_position")
    grid = [[positions[(row, column)] for column in columns] for row in rows]
    headers = [cell.text.strip() for cell in grid[0]]
    if any(
        len(label) < 2
        or len(label) > 40
        or len(label.split()) > 6
        or _shape(label) != "string"
        for label in headers
    ):
        raise _Abstain("header_label_invalid")
    field_ids = [normalize_field_id(label) for label in headers]
    if any(len(field_id) < 2 or len(field_id) > 64 for field_id in field_ids) or len(
        set(field_ids)
    ) != len(field_ids):
        raise _Abstain("header_id_ambiguous")
    for row in grid[1:]:
        if any(
            not cell.text.strip() or len(cell.text) > MAX_CELL_CHARS for cell in row
        ):
            raise _Abstain("blank_or_long_data_cell")
        if [cell.text.strip().casefold() for cell in row] == [
            label.casefold() for label in headers
        ]:
            raise _Abstain("repeated_header")
        if _FOOTER_RE.search(row[0].text.strip()):
            raise _Abstain("footer_or_subtotal")
    body_shapes = [
        [_shape(row[column].text.strip()) for row in grid[1:]]
        for column in range(len(columns))
    ]
    if any(len(set(shapes)) != 1 for shapes in body_shapes):
        raise _Abstain("column_shape_inconsistent")
    if not any(shapes[0] == "number" for shapes in body_shapes) or not any(
        shapes[0] == "string" for shapes in body_shapes
    ):
        raise _Abstain("row_shape_not_data_like")
    return grid, field_ids


def _evidence(
    cell: TableCell, parsed: ParsedDocument, pages: list[dict[str, object]]
) -> tuple[dict[str, object], str]:
    if len(cell.anchors) != 1:
        raise _Abstain("cell_anchor_ambiguous")
    anchor = cell.anchors[0]
    units = {unit.unit_id: index for index, unit in enumerate(parsed.units)}
    if (
        anchor.unit_id not in units
        or not anchor.native_locator
        or anchor.text_span is None
        or anchor.quote != cell.text
    ):
        raise _Abstain("cell_anchor_incomplete")
    page_index = units[anchor.unit_id]
    start, end = anchor.text_span
    if (
        end - start != len(cell.text)
        or pages[page_index]["text"] != parsed.units[page_index].text
        or (parsed.units[page_index].text[start:end] != cell.text)
    ):
        raise _Abstain("cell_span_mismatch")
    _locator(cell)
    return {
        "page": page_index + 1,
        "start": start,
        "end": end,
        "quote": cell.text,
    }, anchor.native_locator


def propose_native_tables(
    parsed: ParsedDocument,
    source_path: Path,
    pages: list[dict[str, object]],
    existing_field_ids: set[str],
    remaining_fields: int,
    *,
    existing_group_ids: set[str] | None = None,
) -> tuple[list[dict[str, object]], list[dict[str, object]], dict[str, object]]:
    """Return v1 linked-column proposals and v2 locator receipts; never release them."""
    report: dict[str, object] = {"tables": [], "cell_provenance": {}}
    fields: list[dict[str, object]] = []
    values: list[dict[str, object]] = []
    if not parsed.tables:
        return fields, values, report
    try:
        if source_path.stat().st_size > MAX_SOURCE_BYTES:
            raise _Abstain("source_too_large")
        source = source_path.read_bytes()
        if len(source) > MAX_SOURCE_BYTES:
            raise _Abstain("source_too_large")
        if (
            hashlib.sha256(source).hexdigest() != parsed.source_sha256
            or source_path.suffix.lower() != f".{parsed.source_format}"
        ):
            raise _Abstain("source_hash_or_format_mismatch")
        if len(pages) != len(parsed.units) or any(
            page.get("page") != index or page.get("text") != unit.text
            for index, (page, unit) in enumerate(zip(pages, parsed.units), start=1)
        ):
            raise _Abstain("source_unit_mismatch")
    except (OSError, _Abstain) as exc:
        report["source_abstention"] = (
            str(exc) if isinstance(exc, _Abstain) else "source_unavailable"
        )
        return fields, values, report

    segments: list[tuple[str, str, list[TableCell], str | None]] = []
    for table in parsed.tables:
        if parsed.source_format == "xlsx":
            by_name: dict[str, list[TableCell]] = defaultdict(list)
            for cell in table.cells:
                try:
                    name = _locator(cell).get("table")
                except _Abstain:
                    name = None
                if isinstance(name, str) and name:
                    by_name[name].append(cell)
            for name, cells in sorted(by_name.items()):
                segments.append(
                    (f"{table.table_id}:{name}", table.table_id, cells, name)
                )
            if not by_name:
                report["tables"].append(
                    {
                        "table_id": table.table_id,
                        "status": "abstained",
                        "reason": "not_declared_xlsx_table",
                    }
                )
        else:
            segments.append((table.table_id, table.table_id, list(table.cells), None))
    if len(segments) > MAX_TABLE_SEGMENTS:
        report["source_abstention"] = "table_segment_limit"
        return fields, values, report

    html_context = None
    workbook = None
    try:
        if parsed.source_format in {"html", "eml"}:
            html_context = _prepare_html(source, parsed.source_format)
        elif parsed.source_format == "xlsx":
            from openpyxl import load_workbook

            workbook = load_workbook(
                BytesIO(source),
                read_only=False,
                data_only=False,
                keep_links=False,
                keep_vba=False,
            )
    except Exception as exc:  # noqa: BLE001 - a source recheck must only abstain, not abort the batch
        report["source_abstention"] = (
            str(exc) if isinstance(exc, _Abstain) else "source_validation_failed"
        )
        return fields, values, report

    eligible: list[tuple[str, str, list[list[TableCell]], list[str], str]] = []
    try:
        for segment_id, parser_table_id, cells, name in segments:
            try:
                grid, field_ids = _rectangular(cells, parsed.source_format)
                if parsed.source_format in {"html", "eml"}:
                    assert html_context is not None
                    header_basis = _validate_html(cells, *html_context)
                elif parsed.source_format == "xlsx":
                    assert name is not None and workbook is not None
                    header_basis = _validate_xlsx(cells, workbook, name)
                elif parsed.source_format in {"docx", "pptx"}:
                    _validate_office_grid_locators(cells, parsed.source_format)
                    header_basis = "content_shape_only"
                else:
                    raise _Abstain("native_format_not_supported_for_rows")
                for cell in cells:
                    _evidence(cell, parsed, pages)
                eligible.append(
                    (segment_id, parser_table_id, grid, field_ids, header_basis)
                )
            except Exception as exc:  # noqa: BLE001 - isolate an untrusted table to an abstention
                reason = (
                    str(exc)
                    if isinstance(exc, _Abstain)
                    else "source_validation_failed"
                )
                report["tables"].append(
                    {
                        "table_id": parser_table_id,
                        "segment_id": segment_id,
                        "status": "abstained",
                        "reason": reason,
                    }
                )
    finally:
        if workbook is not None:
            workbook.close()

    counts = Counter(field_id for _, _, _, ids, _ in eligible for field_id in ids)
    groups = Counter(
        normalize_field_id(f"{grid[0][0].anchors[0].unit_id}_{segment_id}")
        for segment_id, _, grid, _, _ in eligible
    )
    for segment_id, parser_table_id, grid, field_ids, header_basis in eligible:
        if any(
            field_id in existing_field_ids or counts[field_id] > 1
            for field_id in field_ids
        ):
            report["tables"].append(
                {
                    "table_id": parser_table_id,
                    "segment_id": segment_id,
                    "status": "abstained",
                    "reason": "field_id_collision",
                }
            )
            continue
        if len(field_ids) > remaining_fields - len(fields):
            report["tables"].append(
                {
                    "table_id": parser_table_id,
                    "segment_id": segment_id,
                    "status": "abstained",
                    "reason": "schema_budget",
                }
            )
            continue
        group = normalize_field_id(f"{grid[0][0].anchors[0].unit_id}_{segment_id}")
        if not group or groups[group] > 1 or (
            existing_group_ids is not None and group in existing_group_ids
        ):
            report["tables"].append(
                {
                    "table_id": parser_table_id,
                    "segment_id": segment_id,
                    "status": "abstained",
                    "reason": "group_id_collision_or_invalid",
                }
            )
            continue
        provenance = report["cell_provenance"]
        assert isinstance(provenance, dict)
        for column, field_id in enumerate(field_ids):
            fields.append(
                {"field_id": field_id, "type": "array[object]", "group": group}
            )
            linked: list[dict[str, object]] = []
            cell_map: dict[str, object] = {}
            for ordinal, row in enumerate(grid[1:]):
                cell = row[column]
                evidence, locator = _evidence(cell, parsed, pages)
                linked.append(
                    {
                        "row": ordinal,
                        "value": cell.text,
                        "status": "accepted",
                        "evidence": evidence,
                    }
                )
                cell_map[str(ordinal)] = {
                    "source_sha256": parsed.source_sha256,
                    "table_id": parser_table_id,
                    "segment_id": segment_id,
                    "unit_id": cell.anchors[0].unit_id,
                    "grid_row": cell.row,
                    "grid_column": cell.column,
                    "native_locator": locator,
                    "evidence": evidence,
                }
            values.append(
                {
                    "field_id": field_id,
                    "value": linked,
                    "status": "needs_review",
                    "review_state": "pending",
                }
            )
            provenance[field_id] = cell_map
        report["tables"].append(
            {
                "table_id": parser_table_id,
                "segment_id": segment_id,
                "status": "proposed",
                "header_basis": header_basis,
                "field_ids": field_ids,
                "data_rows": len(grid) - 1,
            }
        )
    return fields, values, report
