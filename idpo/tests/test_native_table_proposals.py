"""Review-only native linked-row proposals preserve source-cell identity."""

from __future__ import annotations

import hashlib
import json
from dataclasses import replace
from email.message import EmailMessage
from pathlib import Path

import pytest
from openpyxl import Workbook
from openpyxl.worksheet.table import Table as ExcelTable

from idpo.product_auto.native_formats import parse_native_document
from idpo.product_auto.native_table_proposals import propose_native_tables
from idpo.product_auto.provider import ModelReply
from idpo.product_auto.workflow import process_auto
from product_eval.mixed_format_docling_diagnostic import FIXTURE_NAMES
from product_eval.mixed_format_native_diagnostic import resolve_locator, verify_fixtures

FIXTURES = (
    Path(__file__).resolve().parents[2]
    / "product_eval"
    / "fixtures"
    / "mixed_format_docling_v1"
)
EXPECTED = {
    "item": ["Widget A", "Widget B"],
    "qty": ["2", "1"],
    "amount": ["12.00", "7.50"],
}


class _Model:
    model = "fixture-model"

    def __init__(self, fields: list[dict[str, object]], values: list[dict[str, object]]) -> None:
        self.replies = [
            ModelReply(text=json.dumps({"fields": fields}), latency_ms=1, finish_reason="stop"),
            ModelReply(text=json.dumps({"values": values}), latency_ms=1, finish_reason="stop"),
        ]
        self.calls = 0

    def complete_json(self, *, system: str, prompt: str) -> ModelReply:
        self.calls += 1
        return self.replies.pop(0)


def _model_candidate(
    source: Path,
    tmp_path: Path,
    fields: list[dict[str, object]],
    values: list[dict[str, object]],
) -> tuple[dict, dict, _Model]:
    model = _Model(fields, values)
    outcome = process_auto(
        source, run_id=f"model-table-{source.stem}", model=model,
        output_root=tmp_path / "model-runs",
    )
    run_dir = Path(outcome["run_dir"])
    candidate = json.loads((run_dir / "candidates.jsonl").read_text(encoding="utf-8"))
    sidecar = json.loads(
        (run_dir / "docs" / candidate["doc_id"] / "result.json").read_text(encoding="utf-8")
    )
    return candidate, sidecar, model


def _parsed(path: Path):
    document, parsed = parse_native_document(
        path, hashlib.sha256(path.read_bytes()).hexdigest()
    )
    return parsed, [{"page": page.number, "text": page.text} for page in document.pages]


def _candidate(path: Path, tmp_path: Path) -> tuple[dict, dict]:
    outcome = process_auto(
        path,
        run_id=f"table-{path.stem}-{hashlib.sha256(path.read_bytes()).hexdigest()[:8]}",
        output_root=tmp_path / "runs",
    )
    run_dir = Path(outcome["run_dir"])
    candidate = json.loads((run_dir / "candidates.jsonl").read_text(encoding="utf-8"))
    sidecar = json.loads(
        (run_dir / "docs" / candidate["doc_id"] / "result.json").read_text(
            encoding="utf-8"
        )
    )
    return candidate, sidecar


@pytest.mark.parametrize("name", FIXTURE_NAMES)
def test_five_native_formats_propose_exact_review_only_rows(
    name: str, tmp_path: Path
) -> None:
    manifest = verify_fixtures(FIXTURES)
    source = FIXTURES / name
    candidate, sidecar = _candidate(source, tmp_path)
    assert sidecar["source_sha256"] == manifest["files"][name]
    assert candidate["document_status"] == "needs_review"
    assert sidecar["model_calls"] == []
    if name == "order.docx":
        assert any(
            field["field_id"] == "order_date" and field["type"] == "date"
            for field in candidate["proposed_schema"]
        )
    linked = [
        field
        for field in candidate["proposed_schema"]
        if field["type"] == "array[object]"
    ]
    assert {field["field_id"] for field in linked} == set(EXPECTED)
    assert len({field["group"] for field in linked}) == 1
    parsed_cells = sidecar["parsed_document"]["tables"][0]["cells"]
    cell_map = sidecar["native_table_proposals"]["cell_provenance"]
    for field in linked:
        field_id = field["field_id"]
        proposed = next(
            item
            for item in candidate["extracted_values"]
            if item["field_id"] == field_id
        )
        assert (
            proposed["status"] == "needs_review"
            and proposed["review_state"] == "pending"
        )
        assert [cell["value"] for cell in proposed["value"]] == EXPECTED[field_id]
        assert [cell["row"] for cell in proposed["value"]] == [0, 1]
        for cell in proposed["value"]:
            assert cell["status"] == "accepted"
            evidence = cell["evidence"]
            page_text = sidecar["pages"][evidence["page"] - 1]["text"]
            assert (
                page_text[evidence["start"] : evidence["end"]]
                == evidence["quote"]
                == cell["value"]
            )
            matching = [
                source_cell
                for source_cell in parsed_cells
                if any(
                    anchor["unit_id"]
                    == sidecar["parsed_document"]["units"][evidence["page"] - 1][
                        "unit_id"
                    ]
                    and anchor["text_span"] == [evidence["start"], evidence["end"]]
                    for anchor in source_cell["anchors"]
                )
            ]
            assert len(matching) == 1
            receipt = cell_map[field_id][str(cell["row"])]
            assert receipt["grid_row"] == matching[0]["row"]
            assert receipt["grid_column"] == matching[0]["column"]
            assert receipt["source_sha256"] == manifest["files"][name]
            assert (
                resolve_locator(source, json.loads(receipt["native_locator"]))
                == cell["value"]
            )


@pytest.mark.parametrize("name", ("order.docx", "order.html", "order.xlsx"))
def test_model_values_keep_source_backed_native_rows_for_review(
    name: str, tmp_path: Path,
) -> None:
    manifest = verify_fixtures(FIXTURES)
    source = FIXTURES / name
    candidate, sidecar, model = _model_candidate(
        source, tmp_path,
        [{"field_id": "item_label", "type": "string", "group": "header"}],
        [{"field_id": "item_label", "value": "Widget A", "page": 1, "quote": "Widget A"}],
    )
    assert model.calls == 2
    assert len(sidecar["model_calls"]) == 2
    assert candidate["document_status"] == "needs_review"
    assert sidecar["reason"] == "unapproved_dynamic_schema"
    assert sidecar["source_sha256"] == manifest["files"][name]
    schema = {field["field_id"]: field for field in candidate["proposed_schema"]}
    assert set(schema) == {"item_label", "item", "qty", "amount"}
    assert len({schema[field_id]["group"] for field_id in EXPECTED}) == 1
    values = {value["field_id"]: value for value in candidate["extracted_values"]}
    assert values["item_label"]["value"] == "Widget A"
    assert values["item_label"]["review_state"] == "pending"
    report = sidecar["native_table_proposals"]
    assert [table["status"] for table in report["tables"]] == ["proposed"]
    for field_id, expected in EXPECTED.items():
        assert schema[field_id]["type"] == "array[object]"
        assert values[field_id]["status"] == "needs_review"
        assert values[field_id]["review_state"] == "pending"
        assert [cell["row"] for cell in values[field_id]["value"]] == [0, 1]
        assert [cell["value"] for cell in values[field_id]["value"]] == expected
        for cell in values[field_id]["value"]:
            evidence = cell["evidence"]
            assert sidecar["pages"][evidence["page"] - 1]["text"][
                evidence["start"] : evidence["end"]
            ] == evidence["quote"] == cell["value"]
            receipt = report["cell_provenance"][field_id][str(cell["row"])]
            assert receipt["source_sha256"] == manifest["files"][name]
            assert receipt["evidence"] == evidence
            assert resolve_locator(source, json.loads(receipt["native_locator"])) == cell["value"]


def test_model_field_collision_abstains_whole_native_table(tmp_path: Path) -> None:
    candidate, sidecar, model = _model_candidate(
        FIXTURES / "order.docx", tmp_path,
        [{"field_id": "item", "type": "string"}],
        [{"field_id": "item", "value": "Widget A", "page": 1, "quote": "Widget A"}],
    )
    assert model.calls == 2
    assert [field["field_id"] for field in candidate["proposed_schema"]] == ["item"]
    assert sidecar["native_table_proposals"]["tables"][0]["reason"] == "field_id_collision"
    assert sidecar["native_table_proposals"]["cell_provenance"] == {}


def test_model_group_collision_abstains_whole_native_table(tmp_path: Path) -> None:
    source = FIXTURES / "order.html"
    parsed, pages = _parsed(source)
    proposed, _, _ = propose_native_tables(parsed, source, pages, set(), 24)
    native_group = proposed[0]["group"]
    candidate, sidecar, _ = _model_candidate(
        source, tmp_path,
        [{"field_id": "model_note", "type": "string", "group": native_group.upper()}], [],
    )
    assert [field["field_id"] for field in candidate["proposed_schema"]] == ["model_note"]
    assert sidecar["native_table_proposals"]["tables"][0]["reason"] == "group_id_collision_or_invalid"
    assert sidecar["native_table_proposals"]["cell_provenance"] == {}


def test_model_schema_budget_abstains_without_partial_native_columns(tmp_path: Path) -> None:
    fields = [{"field_id": f"meta_{index}", "type": "string"} for index in range(22)]
    candidate, sidecar, _ = _model_candidate(
        FIXTURES / "order.xlsx", tmp_path, fields, [],
    )
    assert len(candidate["proposed_schema"]) == 22
    assert all(field["type"] == "string" for field in candidate["proposed_schema"])
    assert sidecar["native_table_proposals"]["tables"][0]["reason"] == "schema_budget"
    assert sidecar["native_table_proposals"]["cell_provenance"] == {}


@pytest.mark.parametrize(
    "body",
    [
        (
            "<table><tr><td>Item</td><td>Qty</td><td>Amount</td></tr>"
            "<tr><td>A</td><td>1</td><td>2.00</td></tr><tr><td>B</td><td>2</td><td>3.00</td></tr></table>"
        ),
        (
            '<table role="presentation"><thead><tr><th>Item</th><th>Qty</th><th>Amount</th></tr></thead>'
            "<tr><td>A</td><td>1</td><td>2.00</td></tr><tr><td>B</td><td>2</td><td>3.00</td></tr></table>"
        ),
        (
            "<table><thead><tr><th>Item</th><th>Item</th><th>Amount</th></tr></thead>"
            "<tr><td>A</td><td>1</td><td>2.00</td></tr><tr><td>B</td><td>2</td><td>3.00</td></tr></table>"
        ),
        (
            "<table><thead><tr><th>Item</th><th>Qty</th><th>Amount</th></tr></thead>"
            "<tr><td>A</td><td></td><td>2.00</td></tr><tr><td>B</td><td>2</td><td>3.00</td></tr></table>"
        ),
        (
            "<table><thead><tr><th>Item</th><th>Qty</th><th>Amount</th></tr></thead>"
            "<tr><td>A</td><td>1</td><td>2.00</td></tr><tr><td>Total</td><td>2</td><td>3.00</td></tr></table>"
        ),
        (
            "<table><thead><tr><th>Item</th><th>Qty</th><th>Amount</th></tr></thead>"
            '<tr><td rowspan="2">A</td><td>1</td><td>2.00</td></tr><tr><td>2</td><td>3.00</td></tr></table>'
        ),
    ],
)
def test_html_ambiguous_or_layout_tables_abstain(body: str, tmp_path: Path) -> None:
    source = tmp_path / "adversarial.html"
    source.write_text(f"<html><body>{body}</body></html>", encoding="utf-8")
    candidate, sidecar = _candidate(source, tmp_path)
    assert not any(
        field["type"] == "array[object]" for field in candidate["proposed_schema"]
    )
    assert sidecar["native_table_proposals"]["tables"][0]["status"] == "abstained"


def test_hidden_email_html_table_abstains(tmp_path: Path) -> None:
    message = EmailMessage()
    message.set_content("Fallback body")
    message.add_alternative(
        '<html><body><table aria-hidden="true"><thead><tr><th>Item</th><th>Qty</th>'
        "<th>Amount</th></tr></thead><tr><td>A</td><td>1</td><td>2.00</td></tr>"
        "<tr><td>B</td><td>2</td><td>3.00</td></tr></table></body></html>",
        subtype="html",
    )
    source = tmp_path / "hidden.eml"
    source.write_bytes(message.as_bytes())
    candidate, sidecar = _candidate(source, tmp_path)
    assert not any(
        field["type"] == "array[object]" for field in candidate["proposed_schema"]
    )
    assert not sidecar.get("native_table_proposals", {}).get("tables")


def test_recovered_html5_table_rechecks_visible_cells(tmp_path: Path) -> None:
    source = tmp_path / "modern.html"
    source.write_text(
        '<html><body><header>Current orders</header><table><thead><tr>'
        '<th>Item</th><th>Qty</th><th>Amount</th></tr></thead>'
        '<tbody hidden><tr><td>Hidden</td><td>9</td><td>9.00</td></tr></tbody>'
        '<tbody><tr><td>A<span hidden>SECRET</span></td><td>1</td><td>2.00</td></tr>'
        '<tr><td>B</td><td>2</td><td>3.00</td></tr></tbody></table></body></html>',
        encoding="utf-8",
    )
    candidate, sidecar = _candidate(source, tmp_path)
    linked = [field for field in candidate["proposed_schema"] if field["type"] == "array[object]"]
    assert {field["field_id"] for field in linked} == {"item", "qty", "amount"}
    assert [cell["value"] for cell in next(
        value for value in candidate["extracted_values"] if value["field_id"] == "item"
    )["value"]] == ["A", "B"]
    assert sidecar["native_table_proposals"]["tables"][0]["status"] == "proposed"
    assert "SECRET" not in sidecar["pages"][0]["text"]


def test_declared_windows_1252_email_uses_mime_charset_for_source_recheck(
    tmp_path: Path,
) -> None:
    message = EmailMessage()
    message.set_content("Fallback body")
    message.add_alternative(
        "<html><body><table><thead><tr><th>Item</th><th>Qty</th>"
        "<th>Amount</th></tr></thead><tbody><tr><td>Café A</td><td>1</td>"
        "<td>2.00</td></tr><tr><td>Café B</td><td>2</td><td>3.00</td></tr>"
        "</tbody></table></body></html>",
        subtype="html",
        charset="windows-1252",
    )
    source = tmp_path / "legacy.eml"
    source.write_bytes(message.as_bytes())
    candidate, sidecar = _candidate(source, tmp_path)
    value = next(
        item for item in candidate["extracted_values"] if item["field_id"] == "item"
    )
    assert [cell["value"] for cell in value["value"]] == ["Café A", "Café B"]
    assert sidecar["native_table_proposals"]["tables"][0]["status"] == "proposed"


@pytest.mark.parametrize("variation", ["hidden_row", "formula", "no_table"])
def test_spreadsheet_hidden_formula_or_unstructured_grid_abstains(
    variation: str, tmp_path: Path
) -> None:
    workbook = Workbook()
    sheet = workbook.active
    for row in (("Item", "Qty", "Amount"), ("A", "1", "2.00"), ("B", "2", "3.00")):
        sheet.append(row)
    if variation != "no_table":
        sheet.add_table(ExcelTable(displayName="Items", ref="A1:C3"))
    if variation == "hidden_row":
        sheet.row_dimensions[2].hidden = True
    if variation == "formula":
        sheet["B2"] = "=1"
        sheet["B3"] = "=2"
    source = tmp_path / f"{variation}.xlsx"
    workbook.save(source)
    candidate, sidecar = _candidate(source, tmp_path)
    assert not any(
        field["type"] == "array[object]" for field in candidate["proposed_schema"]
    )
    assert sidecar["native_table_proposals"]["tables"][0]["status"] == "abstained"


def test_word_merged_table_and_two_column_slide_abstain(tmp_path: Path) -> None:
    from docx import Document as WordDocument
    from pptx import Presentation

    word = WordDocument()
    table = word.add_table(rows=3, cols=3)
    for row, values in zip(
        table.rows, (("Item", "Qty", "Amount"), ("A", "1", "2.00"), ("B", "2", "3.00"))
    ):
        for cell, value in zip(row.cells, values):
            cell.text = value
    table.cell(1, 0).merge(table.cell(1, 1))
    docx = tmp_path / "merged.docx"
    word.save(docx)
    docx_candidate, _ = _candidate(docx, tmp_path)
    assert not any(
        field["type"] == "array[object]" for field in docx_candidate["proposed_schema"]
    )

    deck = Presentation()
    slide = deck.slides.add_slide(deck.slide_layouts[5])
    ppt_table = slide.shapes.add_table(3, 2, 0, 0, 2_000_000, 1_000_000).table
    for row_index, values in enumerate((("Item", "Qty"), ("A", "1"), ("B", "2"))):
        for column_index, value in enumerate(values):
            ppt_table.cell(row_index, column_index).text = value
    pptx = tmp_path / "two_columns.pptx"
    deck.save(pptx)
    pptx_candidate, _ = _candidate(pptx, tmp_path)
    assert not any(
        field["type"] == "array[object]" for field in pptx_candidate["proposed_schema"]
    )


def test_hash_change_or_malformed_locator_abstains_without_partial_columns(
    tmp_path: Path,
) -> None:
    source = tmp_path / "order.html"
    source.write_bytes((FIXTURES / "order.html").read_bytes())
    parsed, pages = _parsed(source)
    table = parsed.tables[0]
    first = table.cells[0]
    malformed = replace(
        first, anchors=(replace(first.anchors[0], native_locator="not JSON"),)
    )
    changed = replace(
        parsed, tables=(replace(table, cells=(malformed, *table.cells[1:])),)
    )
    fields, values, report = propose_native_tables(changed, source, pages, set(), 24)
    assert fields == values == []
    assert report["tables"][0]["status"] == "abstained"

    source.write_bytes(source.read_bytes() + b"\n<!-- changed -->")
    fields, values, report = propose_native_tables(parsed, source, pages, set(), 24)
    assert fields == values == []
    assert report["source_abstention"] == "source_hash_or_format_mismatch"


def test_many_small_tables_abstain_before_rechecking_each_table(tmp_path: Path) -> None:
    source = tmp_path / "many.html"
    table = (
        "<table><thead><tr><th>Item</th><th>Qty</th><th>Amount</th></tr></thead>"
        "<tbody><tr><td>A</td><td>1</td><td>2.00</td></tr>"
        "<tr><td>B</td><td>2</td><td>3.00</td></tr></tbody></table>"
    )
    source.write_text(f"<html><body>{table * 33}</body></html>", encoding="utf-8")
    candidate, sidecar = _candidate(source, tmp_path)
    assert not any(
        field["type"] == "array[object]" for field in candidate["proposed_schema"]
    )
    assert (
        sidecar["native_table_proposals"]["source_abstention"] == "table_segment_limit"
    )


def test_table_does_not_overwrite_an_existing_scalar_or_exceed_schema_budget(
    tmp_path: Path,
) -> None:
    source = tmp_path / "collision.html"
    source.write_text(
        "<html><body><p>Item: A</p><table><thead><tr><th>Item</th><th>Qty</th>"
        "<th>Amount</th></tr></thead><tbody><tr><td>A</td><td>1</td><td>2.00</td></tr>"
        "<tr><td>B</td><td>2</td><td>3.00</td></tr></tbody></table></body></html>",
        encoding="utf-8",
    )
    candidate, sidecar = _candidate(source, tmp_path)
    assert any(
        field["field_id"] == "item" and field["type"] == "string"
        for field in candidate["proposed_schema"]
    )
    assert not any(
        field["type"] == "array[object]" for field in candidate["proposed_schema"]
    )
    assert (
        sidecar["native_table_proposals"]["tables"][0]["reason"] == "field_id_collision"
    )

    parsed, pages = _parsed(FIXTURES / "order.html")
    fields, values, report = propose_native_tables(
        parsed, FIXTURES / "order.html", pages, set(), 2
    )
    assert fields == values == []
    assert report["tables"][0]["reason"] == "schema_budget"


def test_duplicate_cell_text_has_distinct_row_and_locator(tmp_path: Path) -> None:
    source = tmp_path / "duplicate_values.html"
    source.write_text(
        "<html><body><table><thead><tr><th>Item</th><th>Qty</th><th>Amount</th>"
        "</tr></thead><tbody><tr><td>A</td><td>1</td><td>2.00</td></tr>"
        "<tr><td>A</td><td>1</td><td>3.00</td></tr></tbody></table></body></html>",
        encoding="utf-8",
    )
    candidate, sidecar = _candidate(source, tmp_path)
    item = next(
        field for field in candidate["extracted_values"] if field["field_id"] == "item"
    )
    assert [cell["value"] for cell in item["value"]] == ["A", "A"]
    assert item["value"][0]["evidence"] != item["value"][1]["evidence"]
    locators = sidecar["native_table_proposals"]["cell_provenance"]["item"]
    assert locators["0"]["native_locator"] != locators["1"]["native_locator"]


def test_supplementary_unicode_keeps_exact_cell_spans(tmp_path: Path) -> None:
    source = tmp_path / "unicode.html"
    source.write_text(
        '<html><head><meta charset="utf-8"></head><body><p>😀 intro</p>'
        "<table><thead><tr><th>Item</th><th>Qty</th>"
        "<th>Amount</th></tr></thead><tbody><tr><td>😀 A</td><td>1</td><td>2.00</td></tr>"
        "<tr><td>😀 B</td><td>2</td><td>3.00</td></tr></tbody></table></body></html>",
        encoding="utf-8",
    )
    candidate, sidecar = _candidate(source, tmp_path)
    value = next(
        item for item in candidate["extracted_values"] if item["field_id"] == "item"
    )
    assert [cell["value"] for cell in value["value"]] == ["😀 A", "😀 B"]
    for cell in value["value"]:
        evidence = cell["evidence"]
        assert (
            sidecar["pages"][evidence["page"] - 1]["text"][
                evidence["start"] : evidence["end"]
            ]
            == cell["value"]
        )
