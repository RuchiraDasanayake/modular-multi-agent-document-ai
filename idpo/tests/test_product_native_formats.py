"""Product intake checks for source-native Office, web and email files."""

from __future__ import annotations

import hashlib
import json
import stat
import warnings
import zipfile
from email.message import EmailMessage
from pathlib import Path

import pytest
from openpyxl import Workbook

from idpo.product_auto.native_formats import NativeLimitError, parse_native_document
from idpo.product_auto.provider import ModelReply
from idpo.product_auto.workflow import process_auto
from product_eval.mixed_format_docling_diagnostic import FIXTURE_NAMES, HEADERS, ROWS
from product_eval.mixed_format_native_diagnostic import resolve_locator, verify_fixtures

FIXTURES = Path(__file__).resolve().parents[2] / "product_eval" / "fixtures" / "mixed_format_docling_v1"
EXPECTED = [value for row in (HEADERS, *ROWS) for value in row]


class _FixedModel:
    model = "source-locked-test"

    def __init__(self) -> None:
        self.replies = [
            '{"fields":[{"field_id":"item_name","type":"string"}]}',
            '{"values":[{"field_id":"item_name","value":"Widget B","quote":"Widget B","page":1}]}',
        ]

    def complete_json(self, *, system: str, prompt: str) -> ModelReply:
        return ModelReply(self.replies.pop(0), 0.1, finish_reason="stop")


@pytest.mark.parametrize("name", FIXTURE_NAMES)
def test_native_source_cell_locators_and_exact_product_spans(name: str, tmp_path: Path) -> None:
    manifest = verify_fixtures(FIXTURES)
    source = FIXTURES / name
    source_hash = hashlib.sha256(source.read_bytes()).hexdigest()
    assert source_hash == manifest["files"][name]
    document, parsed = parse_native_document(source, source_hash)
    assert len(parsed.tables) == 1
    cells = parsed.tables[0].cells
    assert [cell.text for cell in cells] == EXPECTED
    assert all(cell.anchors[0].raw_box is None and cell.anchors[0].display_box is None for cell in cells)
    for cell in cells:
        anchor = cell.anchors[0]
        assert anchor.native_locator is not None
        assert resolve_locator(source, json.loads(anchor.native_locator)) == cell.text
        assert anchor.text_span is not None
        unit = next(unit for unit in parsed.units if unit.unit_id == anchor.unit_id)
        assert unit.text[slice(*anchor.text_span)] == cell.text
    assert [page.text for page in document.pages] == [unit.text for unit in parsed.units]

    outcome = process_auto(source, run_id=f"native-{source.suffix[1:]}", model=_FixedModel(),
                           output_root=tmp_path / "runs")
    assert outcome["status_counts"]["needs_review"] == 1
    run_dir = Path(outcome["run_dir"])
    candidate = json.loads((run_dir / "candidates.jsonl").read_text())
    sidecar = json.loads((run_dir / "docs" / candidate["doc_id"] / "result.json").read_text())
    assert sidecar["parsed_document"]["contract_version"] == 2
    assert sidecar["parsed_document"]["source_sha256"] == source_hash
    assert [page["text"] for page in sidecar["pages"]] == [unit.text for unit in parsed.units]
    value = candidate["extracted_values"][0]
    assert value["value"] == "Widget B" and value["status"] == "accepted"
    assert value["review_state"] == "pending"
    evidence = value["evidence"]
    assert sidecar["pages"][evidence["page"] - 1]["text"][evidence["start"]:evidence["end"]] == "Widget B"
    matching = [cell for cell in cells if cell.text == "Widget B"]
    assert len(matching) == 1
    assert matching[0].anchors[0].text_span == (evidence["start"], evidence["end"])
    assert sidecar["ocr_pages"][0]["engine"] == "native_source"


def _single_result(path: Path, tmp_path: Path) -> tuple[dict, dict]:
    outcome = process_auto(path, run_id="bad-native", output_root=tmp_path / "runs")
    run_dir = Path(outcome["run_dir"])
    candidate = json.loads((run_dir / "candidates.jsonl").read_text())
    sidecar = json.loads((run_dir / "docs" / candidate["doc_id"] / "result.json").read_text())
    return candidate, sidecar


def test_malformed_package_fails_without_aborting_batch(tmp_path: Path) -> None:
    source = tmp_path / "bad.docx"
    source.write_bytes(b"not a ZIP file")
    candidate, sidecar = _single_result(source, tmp_path)
    assert candidate["document_status"] == "failed"
    assert sidecar["reason"] == "parse_failed"


def test_ooxml_rejects_unsafe_member_paths_before_library_parse(tmp_path: Path) -> None:
    source = tmp_path / "unsafe.docx"
    with zipfile.ZipFile(source, "w") as archive:
        archive.writestr("[Content_Types].xml", "<Types/>")
        archive.writestr("word/document.xml", "<document/>")
        archive.writestr("../outside.xml", "<x/>")
    candidate, sidecar = _single_result(source, tmp_path)
    assert candidate["document_status"] == "failed"
    assert sidecar["reason"] == "parse_failed"


def test_ooxml_rejects_ambiguous_members_symlinks_and_uppercase_xml_dtd(tmp_path: Path) -> None:
    def package(name: str, member: zipfile.ZipInfo | str, content: bytes) -> Path:
        source = tmp_path / name
        with zipfile.ZipFile(source, "w") as archive:
            archive.writestr("[Content_Types].xml", "<Types/>")
            archive.writestr("word/document.xml", "<document/>")
            archive.writestr(member, content)
        return source

    with warnings.catch_warnings():
        warnings.simplefilter("ignore", UserWarning)
        duplicate = package("duplicate.docx", "word/document.xml", b"<different/>")
    with pytest.raises(ValueError, match="duplicate_archive_member"):
        parse_native_document(duplicate, hashlib.sha256(duplicate.read_bytes()).hexdigest())

    link = zipfile.ZipInfo("word/link.xml")
    link.create_system = 3
    link.external_attr = (stat.S_IFLNK | 0o777) << 16
    symlink = package("symlink.docx", link, b"target")
    with pytest.raises(ValueError, match="unsafe_archive_member"):
        parse_native_document(symlink, hashlib.sha256(symlink.read_bytes()).hexdigest())

    dtd = package("dtd.docx", "word/EXTRA.XML", b"<!doctype x [<!entity x 'bad'>]><x/>")
    with pytest.raises(ValueError, match="ooxml_dtd_forbidden"):
        parse_native_document(dtd, hashlib.sha256(dtd.read_bytes()).hexdigest())


def test_archive_expansion_limit_abstains_before_library_parse(tmp_path: Path, monkeypatch) -> None:
    from idpo.product_auto import native_formats

    monkeypatch.setattr(native_formats, "MAX_ARCHIVE_UNCOMPRESSED", 1)
    candidate, sidecar = _single_result(FIXTURES / "order.docx", tmp_path)
    assert candidate["document_status"] == "unsupported"
    assert sidecar["reason"] == "native_limit_exceeded"
    assert "parsed_document" not in sidecar


def test_sparse_sheet_is_bounded_and_formulas_are_not_evaluated(tmp_path: Path) -> None:
    source = tmp_path / "formula.xlsx"
    book = Workbook()
    sheet = book.active
    sheet["A1"] = "Total"
    sheet["B1"] = "=2+2"
    book.save(source)
    parsed = parse_native_document(source, hashlib.sha256(source.read_bytes()).hexdigest())[1]
    formula = next(cell for cell in parsed.tables[0].cells if cell.column == 1)
    assert formula.text == "=2+2"
    assert json.loads(formula.anchors[0].native_locator)["formula"] is True

    sparse = tmp_path / "sparse.xlsx"
    sheet["Z5000"] = "far away"
    book.save(sparse)
    with pytest.raises(NativeLimitError, match="sheet_grid_limit"):
        parse_native_document(sparse, hashlib.sha256(sparse.read_bytes()).hexdigest())


def test_html_uses_local_text_only_and_rejects_excessive_structure(tmp_path: Path) -> None:
    source = tmp_path / "source.html"
    source.write_text('<html><body><img src="https://example.invalid/remote.png"><script>SECRET</script>'
                      '<table><tr><td rowspan="2">A</td><td>B</td></tr><tr><td>C</td></tr></table>'
                      '</body></html>', encoding="utf-8")
    parsed = parse_native_document(source, hashlib.sha256(source.read_bytes()).hexdigest())[1]
    assert "SECRET" not in parsed.units[0].text
    assert [cell.text for cell in parsed.tables[0].cells] == ["A", "B", "C"]
    assert parsed.tables[0].cells[0].row_span == 2
    assert parsed.tables[0].cells[2].column == 1
    assert all(anchor.raw_box is None for cell in parsed.tables[0].cells for anchor in cell.anchors)

    source.write_text("<html><body>" + "<p>x</p>" * 100 + "</body></html>", encoding="utf-8")
    from idpo.product_auto import native_formats
    original_limit = native_formats.MAX_HTML_ELEMENTS
    native_formats.MAX_HTML_ELEMENTS = 20
    try:
        with pytest.raises(NativeLimitError, match="html_element_limit"):
            parse_native_document(source, hashlib.sha256(source.read_bytes()).hexdigest())
    finally:
        native_formats.MAX_HTML_ELEMENTS = original_limit

    source.write_text("<table><tr><td><table><tr><td>nested</td></tr></table></td></tr></table>", encoding="utf-8")
    with pytest.raises(NativeLimitError, match="nested_html_table_unsupported"):
        parse_native_document(source, hashlib.sha256(source.read_bytes()).hexdigest())


def test_email_selects_html_body_not_attachment_and_rejects_defects(tmp_path: Path) -> None:
    message = EmailMessage()
    message["Subject"] = "Order"
    message.set_content("plain body")
    message.add_alternative("<html><body><table><tr><td>Visible</td></tr></table></body></html>", subtype="html")
    message.add_attachment(b"attachment secret", maintype="application", subtype="octet-stream", filename="x.bin")
    source = tmp_path / "order.eml"
    source.write_bytes(message.as_bytes())
    parsed = parse_native_document(source, hashlib.sha256(source.read_bytes()).hexdigest())[1]
    assert parsed.tables[0].cells[0].text == "Visible"
    assert "attachment secret" not in parsed.units[0].text
    assert json.loads(parsed.tables[0].cells[0].anchors[0].native_locator)["mime_path"] == [0, 1]

    source.write_bytes(b"MIME-Version: 1.0\nContent-Type: multipart/mixed; boundary=absent\n\nno boundary")
    candidate, sidecar = _single_result(source, tmp_path)
    assert candidate["document_status"] == "failed"
    assert sidecar["reason"] == "parse_failed"


def test_email_html_preserves_declared_legacy_charset(tmp_path: Path) -> None:
    message = EmailMessage()
    message.set_content("fallback")
    message.add_alternative('<html><head><meta charset="iso-8859-1"></head><body>'
                            '<table><tr><td>Café</td></tr></table></body></html>',
                            subtype="html", charset="iso-8859-1")
    source = tmp_path / "latin1.eml"
    source.write_bytes(message.as_bytes())
    parsed = parse_native_document(source, hashlib.sha256(source.read_bytes()).hexdigest())[1]
    assert parsed.tables[0].cells[0].text == "Café"


def test_docx_merged_cell_has_one_resolvable_spanning_anchor(tmp_path: Path) -> None:
    from docx import Document as WordDocument

    source = tmp_path / "merged.docx"
    document = WordDocument()
    table = document.add_table(rows=1, cols=2)
    table.cell(0, 0).merge(table.cell(0, 1)).text = "Merged"
    document.save(source)
    parsed = parse_native_document(source, hashlib.sha256(source.read_bytes()).hexdigest())[1]
    assert len(parsed.tables[0].cells) == 1
    cell = parsed.tables[0].cells[0]
    assert cell.column_span == 2 and cell.text == "Merged"
    assert resolve_locator(source, json.loads(cell.anchors[0].native_locator)) == "Merged"

    nested = tmp_path / "nested.docx"
    table.cell(0, 0).add_table(rows=1, cols=1)
    document.save(nested)
    with pytest.raises(NativeLimitError, match="nested_docx_table_unsupported"):
        parse_native_document(nested, hashlib.sha256(nested.read_bytes()).hexdigest())
