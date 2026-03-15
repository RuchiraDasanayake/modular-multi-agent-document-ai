"""Inbox-based document processing with batch orchestration and result export."""

from __future__ import annotations

import hashlib
import json
import logging
import mimetypes
import os
import re
import shutil
import time
import uuid
from pathlib import Path
from typing import Any, Iterable

from idpo.agents.intake import IntakeAgent
from idpo.core.config import PipelineConfig, load_config
from idpo.core.preflight import run_preflight
from idpo.pipeline.orchestrator import Pipeline

SUPPORTED_INPUT_EXTS = {".json", ".txt", ".pdf", ".png", ".jpg", ".jpeg", ".tif", ".tiff", ".bmp"}
logger = logging.getLogger("idpo.inbox")

BASE_DOC_HEADERS = [
    "doc_id",
    "source",
    "doc_type",
    "vendor_cluster",
    "quality_bin",
    "complexity_bin",
    "content_profile",
    "quality_score",
    "complexity_score",
    "policy_id",
    "policy_name",
    "plan_id",
    "plan_source",
    "model_used",
    "model_mode",
    "model_reason",
    "latency_ms",
    "preprocess_route",
    "ocr_engine_used",
    "ocr_avg_confidence",
    "ocr_rerun_attempted_pages",
    "ocr_rerun_used_pages",
    "fallback_path",
    "unresolved_field_count",
    "entity_count",
    "explicit_entity_count",
    "default_entity_count",
    "validation_issue_count",
    "recovered_fields_count",
    "requested_field_status",
    "llm_diagnostics",
    "escalated",
    "chunk_count",
    "summary",
    "input_original_path",
    "input_manifest_path",
    "input_hashes_path",
    "dae_output_path",
    "plan_path",
    "ocr_path",
    "structuring_extracted_path",
    "structuring_sections_path",
    "structuring_summary_path",
    "validation_results_path",
    "validation_normalized_path",
    "recovery_attempts_path",
    "recovery_escalations_path",
    "chunking_path",
    "embeddings_index_path",
    "provenance_graph_path",
    "provenance_events_path",
    "transparency_path",
    "artifacts_path",
]

def process_inbox(
    root: Path,
    config_path: str | None = None,
    extraction_fields: list[dict[str, str]] | None = None,
    run_id: str | None = None,
    file_paths: list[str] | None = None,
) -> dict[str, Any]:
    config = load_config(config_path)
    config.shadow.enabled = False
    preflight_warnings = run_preflight(config)

    root = Path(root)
    inbox_dir = root / "inputs" / "inbox"
    processed_dir = root / "inputs" / "processed"
    runs_dir = root / "runs"
    inbox_dir.mkdir(parents=True, exist_ok=True)
    processed_dir.mkdir(parents=True, exist_ok=True)
    runs_dir.mkdir(parents=True, exist_ok=True)

    if file_paths:
        files = [Path(p) for p in file_paths if Path(p).exists() and Path(p).suffix.lower() in SUPPORTED_INPUT_EXTS]
    else:
        files = [
            path for path in inbox_dir.iterdir() if path.is_file() and path.suffix.lower() in SUPPORTED_INPUT_EXTS
        ]
    if not files:
        raise ValueError(f"No supported files found in {inbox_dir}. Drop files there and retry.")

    resolved_run_id = run_id or uuid.uuid4().hex
    run_dir = runs_dir / resolved_run_id
    docs_dir = run_dir / "docs"
    docs_dir.mkdir(parents=True, exist_ok=True)

    expected_fields = _normalize_fields(extraction_fields or [])
    if expected_fields:
        _write_fields_schema(run_dir, expected_fields)
        _apply_field_overrides(config, run_dir, expected_fields)

    events_path = run_dir / "events.jsonl"
    errors_path = run_dir / "errors.jsonl"
    start_time = time.time()
    run_payload = {
        "run_id": resolved_run_id,
        "started_at": start_time,
        "config_path": config_path,
        "policy_store_path": config.knowledge.policy_store_path,
        "preflight_warnings": " | ".join(preflight_warnings),
    }
    (run_dir / "run.json").write_text(json.dumps(run_payload, indent=2), encoding="utf-8")

    intake = IntakeAgent()
    failed: list[dict[str, str]] = []
    doc_records: list[dict[str, object]] = []
    ingested_items: list[dict[str, Any]] = []
    ingested_files: list[Path] = []

    for file_path in files:
        try:
            docs = intake.ingest(str(file_path))
            if not docs:
                raise ValueError("No documents were parsed from the input")
            for doc in docs:
                doc_dir = docs_dir / _safe_path_component(doc.doc_id)
                _ensure_doc_dirs(doc_dir)
                _copy_original_input(file_path, doc_dir)
                _write_doc_manifest(doc, doc_dir)
                _write_hashes(file_path, doc, doc_dir)
                ingested_items.append({"doc": doc, "doc_dir": doc_dir, "file_path": file_path})
            ingested_files.append(file_path)
        except Exception as exc:
            error_msg = str(exc)
            failed.append({"file": file_path.name, "error": error_msg})
            errors_path.parent.mkdir(parents=True, exist_ok=True)
            with errors_path.open("a", encoding="utf-8") as handle:
                handle.write(json.dumps({"file": str(file_path), "error": error_msg}) + "\n")
            failed_dir = root / "inputs" / "failed"
            failed_dir.mkdir(parents=True, exist_ok=True)
            try:
                shutil.move(str(file_path), str(failed_dir / file_path.name))
            except Exception:
                pass

    if not ingested_items:
        raise ValueError("No valid documents were ingested for processing.")

    # Throughput mode: process all docs in a single pipeline run and auto-enable queueing for multi-doc batches.
    if len(ingested_items) > 1 and not config.queue.enabled:
        config.queue.enabled = True
        config.queue.max_workers = max(1, min(config.queue.max_workers, len(ingested_items)))
        preflight_warnings.append("Queue mode auto-enabled for multi-document batch processing")

    pipeline = Pipeline(config=config, run_dir=str(run_dir))
    docs_batch = [item["doc"] for item in ingested_items]
    results: list[Any] = []
    failed_doc_ids: set[str] = set()

    try:
        results = pipeline.process_documents(docs_batch)
    except Exception as exc:
        # Preserve progress by falling back to per-document execution in the same pipeline run.
        logger.warning("Batch processing failed, falling back to per-document mode: %s", exc)
        preflight_warnings.append(f"Batch mode fallback triggered: {str(exc)[:180]}")
        for item in ingested_items:
            doc = item["doc"]
            try:
                results.append(pipeline.process_document(doc))
            except Exception as inner_exc:
                failed_doc_ids.add(str(doc.doc_id))
                error_msg = str(inner_exc)
                failed.append({"file": Path(item["file_path"]).name, "error": f"{doc.doc_id}: {error_msg}"})
                with errors_path.open("a", encoding="utf-8") as handle:
                    handle.write(
                        json.dumps(
                            {"file": str(item["file_path"]), "doc_id": doc.doc_id, "error": error_msg}
                        )
                        + "\n"
                    )

    if results:
        pipeline.build_index(results, directory=str(run_dir / "embeddings" / "index"))

    results_by_doc_id = {str(result.doc_id): result for result in results}
    for item in ingested_items:
        doc = item["doc"]
        doc_dir = item["doc_dir"]
        result = results_by_doc_id.get(str(doc.doc_id))
        if not result:
            failed_doc_ids.add(str(doc.doc_id))
            continue
        pipeline.build_index([result], directory=str(doc_dir / "embeddings" / "index"))
        _write_dae_output(result, doc_dir)
        _write_plan_output(result, doc_dir)
        _write_structuring(result, doc_dir)
        _write_validation(result, doc_dir)
        _write_recovery(result, doc_dir)
        _write_chunking(result, doc_dir)
        _write_ocr_output(doc, result, doc_dir)
        _write_transparency(result, doc_dir)
        _materialize_doc_provenance_from_run(run_dir, doc_dir, str(doc.doc_id))
        _append_events(doc_dir, events_path)
        _normalize_doc_layout(doc_dir)
        doc_records.append(_doc_record(result, doc_dir, expected_fields, doc))

    for file_path in ingested_files:
        if not file_path.exists():
            continue
        try:
            shutil.move(str(file_path), str(processed_dir / file_path.name))
        except Exception:
            pass

    metrics_payload = {
        "file_count": len(files),
        "doc_count": len(ingested_items),
        "processed_doc_count": len(doc_records),
        "failed_count": len(failed),
        "failed_doc_count": len(failed_doc_ids),
        "ended_at": time.time(),
        "duration_ms": (time.time() - start_time) * 1000,
    }
    (run_dir / "metrics.json").write_text(json.dumps(metrics_payload, indent=2), encoding="utf-8")
    _write_docs_excel(run_dir, doc_records)
    _write_entities_excel(run_dir, doc_records, expected_fields)
    _write_run_excel(run_dir, run_payload, metrics_payload)
    telemetry_status = _persist_run_telemetry(run_dir, resolved_run_id, doc_records, metrics_payload, preflight_warnings)
    if telemetry_status and telemetry_status.get("warning"):
        preflight_warnings.append(str(telemetry_status["warning"]))
    run_payload["telemetry"] = telemetry_status
    (run_dir / "run.json").write_text(json.dumps(run_payload, indent=2), encoding="utf-8")

    return {
        "run_id": resolved_run_id,
        "run_dir": str(run_dir),
        "doc_records": doc_records,
        "entity_records": _entity_records(doc_records, expected_fields),
        "metrics": metrics_payload,
        "run_payload": run_payload,
        "preflight_warnings": preflight_warnings,
        "failed": failed,
    }

def _apply_field_overrides(config: PipelineConfig, run_dir: Path, fields: list[dict[str, str]]) -> None:
    names = [field["key"] for field in fields]
    descriptions = {field["key"]: field.get("description", "") for field in fields if field.get("description")}
    override = {
        "default_policy_id": "custom_fields",
        "policies": [
            {
                "policy_id": "custom_fields",
                "name": "Custom Fields",
                "conditions": {},
                "actions": {
                    "execution_mode": "extraction_only",
                    "extraction": {
                        "explicit_entities": names,
                        "field_descriptions": descriptions,
                    }
                },
            }
        ],
    }
    policy_path = run_dir / "policy_override.json"
    policy_path.write_text(json.dumps(override, indent=2), encoding="utf-8")
    config.knowledge.policy_store_path = str(policy_path)
    config.llm.use_extraction = True

def _normalize_fields(fields: list[dict[str, str]]) -> list[dict[str, str]]:
    normalized = []
    for field in fields:
        name = str(field.get("name", "")).strip()
        description = str(field.get("description", "")).strip()
        if not name:
            continue
        key = _normalize_key(name)
        normalized.append({"name": name, "key": key, "description": description})
    return normalized

def _write_fields_schema(run_dir: Path, fields: list[dict[str, str]]) -> None:
    schema = {"fields": fields}
    (run_dir / "fields.json").write_text(json.dumps(schema, indent=2), encoding="utf-8")

def _normalize_key(raw: str) -> str:
    return "".join(ch.lower() if ch.isalnum() else "_" for ch in raw).strip("_")

def _safe_path_component(raw: object) -> str:
    value = re.sub(r"[^a-zA-Z0-9._-]+", "_", str(raw or "")).strip("._-")
    return value[:120] if value else uuid.uuid4().hex

def _ensure_doc_dirs(doc_dir: Path) -> None:
    for folder in [
        doc_dir / "input",
        doc_dir / "dae",
        doc_dir / "plan" / "shadow",
        doc_dir / "preprocessing" / "ocr",
        doc_dir / "preprocessing" / "layout",
        doc_dir / "structuring",
        doc_dir / "validation",
        doc_dir / "recovery",
        doc_dir / "chunking",
        doc_dir / "embeddings",
        doc_dir / "retrieval",
        doc_dir / "qa",
        doc_dir / "provenance",
    ]:
        folder.mkdir(parents=True, exist_ok=True)

def _copy_original_input(file_path: Path, doc_dir: Path) -> None:
    if not file_path.exists():
        return
    dest = doc_dir / "input" / f"original{file_path.suffix.lower()}"
    shutil.copy2(file_path, dest)

def _write_doc_manifest(doc, doc_dir: Path) -> None:
    payload = {
        "doc_id": doc.doc_id,
        "source": doc.source,
        "path": doc.path,
        "metadata": doc.metadata,
    }
    (doc_dir / "input" / "manifest.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")

def _write_hashes(file_path: Path, doc, doc_dir: Path) -> None:
    sha256 = hashlib.sha256()
    with file_path.open("rb") as handle:
        for chunk in iter(lambda: handle.read(8192), b""):
            sha256.update(chunk)
    mime = mimetypes.guess_type(str(file_path))[0] or "application/octet-stream"
    payload = {
        "sha256": sha256.hexdigest(),
        "size_bytes": file_path.stat().st_size,
        "mime_type": mime,
        "pages": len(getattr(doc, "pages", []) or []),
    }
    (doc_dir / "input" / "hashes.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")

def _write_dae_output(result, doc_dir: Path) -> None:
    payload = {
        "doc_type": result.doc_type,
        "vendor_cluster": result.vendor_cluster,
        "quality_bin": result.quality_bin,
        "complexity_bin": result.complexity_bin,
        "content_profile": result.content_profile,
        "quality_score": result.quality_score,
        "complexity_score": result.complexity_score,
        "signals": result.dae_signals,
    }
    (doc_dir / "dae" / "output.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")

def _write_plan_output(result, doc_dir: Path) -> None:
    payload = result.plan_payload or {}
    (doc_dir / "plan" / "primary.json").write_text(json.dumps(payload, indent=2), encoding="utf-8")

def _write_structuring(result, doc_dir: Path) -> None:
    requested_status = getattr(result, "field_details", {}) or {}
    unresolved_fields = sorted(
        key
        for key, item in requested_status.items()
        if str((item or {}).get("status", "")).strip() != "extracted"
    )
    extracted = {
        "entities": [entity.__dict__ for entity in result.entities],
        "explicit_entities": [entity.__dict__ for entity in result.explicit_entities],
        "default_entities": [entity.__dict__ for entity in result.default_entities],
        "explicit_source": result.explicit_source,
        "requested_field_status": requested_status,
        "unresolved_fields": unresolved_fields,
        "llm_diagnostics": getattr(result, "llm_diagnostics", {}) or {},
    }
    (doc_dir / "structuring" / "extracted.json").write_text(json.dumps(extracted, indent=2), encoding="utf-8")
    sections = {"sections": [section.__dict__ for section in result.sections]}
    (doc_dir / "structuring" / "sections.json").write_text(json.dumps(sections, indent=2), encoding="utf-8")
    (doc_dir / "structuring" / "summaries.json").write_text(
        json.dumps({"summary": result.summary}, indent=2),
        encoding="utf-8",
    )

def _write_validation(result, doc_dir: Path) -> None:
    (doc_dir / "validation" / "results.json").write_text(
        json.dumps({"issues": [issue.__dict__ for issue in result.validation_issues]}, indent=2),
        encoding="utf-8",
    )
    (doc_dir / "validation" / "normalized.json").write_text(
        json.dumps(result.normalized_entities, indent=2),
        encoding="utf-8",
    )

def _write_recovery(result, doc_dir: Path) -> None:
    (doc_dir / "recovery" / "attempts.jsonl").write_text(
        json.dumps(result.recovered_fields, indent=2),
        encoding="utf-8",
    )
    (doc_dir / "recovery" / "escalations.jsonl").write_text(
        json.dumps({"escalated": result.escalated}, indent=2),
        encoding="utf-8",
    )

def _write_chunking(result, doc_dir: Path) -> None:
    chunks_path = doc_dir / "chunking" / "chunks.jsonl"
    with chunks_path.open("w", encoding="utf-8") as handle:
        for chunk in result.chunks:
            handle.write(json.dumps(chunk.__dict__) + "\n")

def _write_provenance(doc_dir: Path) -> None:
    graph_path = doc_dir / "provenance_graph.json"
    if not graph_path.exists():
        graph_path = doc_dir / "provenance" / "provenance_graph.json"
    if graph_path.exists():
        target = doc_dir / "provenance" / "graph.json"
        shutil.copy2(graph_path, target)

def _materialize_doc_provenance_from_run(run_dir: Path, doc_dir: Path, doc_id: str) -> None:
    provenance_dir = doc_dir / "provenance"
    provenance_dir.mkdir(parents=True, exist_ok=True)
    events_src = run_dir / "provenance" / "events.jsonl"
    graph_src = run_dir / "provenance" / "provenance_graph.json"

    if events_src.exists():
        target_events = provenance_dir / "events.jsonl"
        with events_src.open("r", encoding="utf-8") as src, target_events.open("w", encoding="utf-8") as dst:
            for line in src:
                stripped = line.strip()
                if not stripped:
                    continue
                try:
                    payload = json.loads(stripped)
                except Exception:
                    continue
                if str((payload.get("inputs") or {}).get("doc_id", "")) == doc_id:
                    dst.write(stripped + "\n")

    if graph_src.exists():
        shutil.copy2(graph_src, provenance_dir / "graph.json")

    for name in ("run_metadata.json", "policy_rollbacks.json"):
        src = run_dir / name
        if src.exists():
            shutil.copy2(src, provenance_dir / name)

    # Keep expected directory contract for downstream consumers.
    (provenance_dir / "artifacts").mkdir(parents=True, exist_ok=True)

def _write_ocr_output(doc, result, doc_dir: Path) -> None:
    ocr_signal = None
    for signal in result.stage_signals:
        if signal.stage == "ocr":
            ocr_signal = signal
            break
    pages = []
    for page in doc.pages:
        pages.append(
            {
                "number": page.number,
                "text": page.text,
                "confidence": ocr_signal.confidence if ocr_signal else None,
            }
        )
    payload = {"pages": pages, "confidence": ocr_signal.confidence if ocr_signal else None}
    (doc_dir / "preprocessing" / "ocr" / "text.json").write_text(
        json.dumps(payload, indent=2),
        encoding="utf-8",
    )

def _write_transparency(result, doc_dir: Path) -> None:
    ocr_signal = None
    for signal in result.stage_signals:
        if signal.stage == "ocr":
            ocr_signal = signal
            break
    ocr_meta = getattr(ocr_signal, "metadata", {}) or {}
    llm_diag = getattr(result, "llm_diagnostics", {}) or {}
    requested_status = getattr(result, "field_details", {}) or {}
    unresolved = sorted(
        key
        for key, item in requested_status.items()
        if str((item or {}).get("status", "")).strip() != "extracted"
    )
    payload = {
        "doc_id": result.doc_id,
        "ocr_route": {
            "engine": str(ocr_meta.get("engine", "")),
            "avg_confidence": float(ocr_meta.get("avg_ocr_confidence", 0.0) or 0.0),
            "rerun_attempted_pages": int(ocr_meta.get("rerun_attempted_pages", 0) or 0),
            "rerun_used_pages": int(ocr_meta.get("rerun_used_pages", 0) or 0),
        },
        "model_route": {
            "model_used": getattr(result, "model_used", ""),
            "model_mode": getattr(result, "model_mode", ""),
            "model_reason": getattr(result, "model_reason", ""),
            "fallback_path": _fallback_path(llm_diag if isinstance(llm_diag, dict) else {}),
        },
        "field_resolution": {
            "requested_field_status": requested_status,
            "unresolved_fields": unresolved,
            "unresolved_count": len(unresolved),
        },
        "latency_ms": float(getattr(result, "latency_ms", 0.0) or 0.0),
    }
    (doc_dir / "provenance" / "transparency.json").write_text(
        json.dumps(payload, indent=2),
        encoding="utf-8",
    )

def _append_events(doc_dir: Path, events_path: Path) -> None:
    src = doc_dir / "provenance.jsonl"
    if not src.exists():
        src = doc_dir / "provenance" / "events.jsonl"
    if not src.exists():
        return
    with src.open("r", encoding="utf-8") as handle, events_path.open("a", encoding="utf-8") as out:
        for line in handle:
            if line.strip():
                out.write(line)

def _normalize_doc_layout(doc_dir: Path) -> None:
    provenance_dir = doc_dir / "provenance"
    provenance_dir.mkdir(parents=True, exist_ok=True)
    moves = {
        doc_dir / "provenance_graph.json": provenance_dir / "graph.json",
        doc_dir / "provenance.jsonl": provenance_dir / "events.jsonl",
        provenance_dir / "provenance_graph.json": provenance_dir / "graph.json",
        doc_dir / "run_metadata.json": provenance_dir / "run_metadata.json",
        doc_dir / "policy_rollbacks.json": provenance_dir / "policy_rollbacks.json",
    }
    for src, dest in moves.items():
        if src.exists():
            try:
                if src.resolve() == dest.resolve():
                    continue
            except Exception:
                pass
            if dest.exists():
                dest.unlink()
            shutil.move(str(src), str(dest))
    artifacts_dir = doc_dir / "artifacts"
    if artifacts_dir.exists():
        shutil.move(str(artifacts_dir), str(provenance_dir / "artifacts"))

def _doc_record(result, doc_dir: Path, expected_fields: list[dict[str, str]], doc=None) -> dict[str, object]:
    # Build extracted full text from document pages
    extracted_text = ""
    if doc and hasattr(doc, "pages") and doc.pages:
        extracted_text = "\n\n".join(
            f"--- Page {page.number} ---\n{page.text}"
            for page in doc.pages
            if page.text
        )

    # Build structured entities dict for the frontend
    entities_dict: dict[str, object] = {}
    entity_details: list[dict[str, object]] = []
    for entity in result.entities:
        key = _normalize_key(entity.key)
        entities_dict[key] = entity.value
        entity_details.append({
            "key": key,
            "label": entity.key,
            "value": entity.value,
            "confidence": round(entity.confidence, 3) if entity.confidence is not None else 0,
            "page": entity.page,
            "source": (entity.metadata or {}).get("source", "unknown"),
            "is_explicit": bool((entity.metadata or {}).get("explicit")),
        })

    expected_keys = [field["key"] for field in expected_fields]
    requested_field_status = getattr(result, "field_details", {}) or {}
    llm_diagnostics = getattr(result, "llm_diagnostics", {}) or {}
    ocr_signal = None
    for signal in getattr(result, "stage_signals", []) or []:
        if getattr(signal, "stage", "") == "ocr":
            ocr_signal = signal
            break
    ocr_meta = getattr(ocr_signal, "metadata", {}) or {}
    fallback_path = _fallback_path(llm_diagnostics) if isinstance(llm_diagnostics, dict) else "none"
    unresolved_field_count = sum(
        1 for item in requested_field_status.values() if str((item or {}).get("status", "")).strip() != "extracted"
    )
    for key in expected_keys:
        entities_dict.setdefault(key, "")

    # Build sections list for the frontend
    sections_list = [
        {"name": s.name, "text": s.text[:500], "confidence": round(s.confidence, 3)}
        for s in (result.sections or [])
    ]

    # Build validation issues for the frontend
    validation_list = [
        {"field": vi.field, "issue": vi.issue, "severity": vi.severity}
        for vi in (result.validation_issues or [])
    ]

    # Determine original filename from path
    original_path = Path(result.path)
    filename = original_path.name

    record = {
        "doc_id": result.doc_id,
        "filename": filename,
        "source": result.source,
        "status": "success",
        "category": result.doc_type or result.category or "",
        "doc_type": result.doc_type,
        "vendor_cluster": result.vendor_cluster,
        "quality_bin": result.quality_bin,
        "complexity_bin": result.complexity_bin,
        "content_profile": result.content_profile,
        "quality_score": result.quality_score,
        "complexity_score": result.complexity_score,
        "policy_id": result.policy_id,
        "policy_name": result.policy_name,
        "plan_id": result.plan_id,
        "plan_source": result.plan_source,
        "model_used": getattr(result, "model_used", ""),
        "model_mode": getattr(result, "model_mode", ""),
        "model_reason": getattr(result, "model_reason", ""),
        "latency_ms": round(float(getattr(result, "latency_ms", 0.0) or 0.0), 2),
        "preprocess_route": ",".join(result.preprocess_route),
        "ocr_engine_used": str(ocr_meta.get("engine", "")),
        "ocr_avg_confidence": round(float(ocr_meta.get("avg_ocr_confidence", 0.0) or 0.0), 2),
        "ocr_rerun_attempted_pages": int(ocr_meta.get("rerun_attempted_pages", 0) or 0),
        "ocr_rerun_used_pages": int(ocr_meta.get("rerun_used_pages", 0) or 0),
        "fallback_path": fallback_path,
        "unresolved_field_count": unresolved_field_count,
        "entity_count": len(result.entities),
        "explicit_entity_count": len(result.explicit_entities),
        "default_entity_count": len(result.default_entities),
        "validation_issue_count": len(result.validation_issues),
        "recovered_fields_count": len(result.recovered_fields),
        "escalated": result.escalated,
        "chunk_count": len(result.chunks),
        "summary": result.summary,
        "extracted_text": extracted_text,
        "entities": entities_dict,
        "entity_details": entity_details,
        "sections": sections_list,
        "validation_issues": validation_list,
        "field_confidence": result.field_confidence or {},
        "recovered_fields": result.recovered_fields or {},
        "requested_field_status": requested_field_status,
        "llm_diagnostics": llm_diagnostics,
        "structured_fields": {},
        "missing_fields": [],
        "transparency": {
            "ocr_route": {
                "engine": str(ocr_meta.get("engine", "")),
                "avg_confidence": float(ocr_meta.get("avg_ocr_confidence", 0.0) or 0.0),
                "rerun_attempted_pages": int(ocr_meta.get("rerun_attempted_pages", 0) or 0),
                "rerun_used_pages": int(ocr_meta.get("rerun_used_pages", 0) or 0),
            },
            "model_route": {
                "model_used": getattr(result, "model_used", ""),
                "model_mode": getattr(result, "model_mode", ""),
                "model_reason": getattr(result, "model_reason", ""),
                "fallback_path": fallback_path,
            },
            "field_resolution": {
                "requested": len(requested_field_status),
                "unresolved": unresolved_field_count,
            },
        },
        "input_original_path": _find_original_path(doc_dir),
        "input_manifest_path": str(doc_dir / "input" / "manifest.json"),
        "input_hashes_path": str(doc_dir / "input" / "hashes.json"),
        "dae_output_path": str(doc_dir / "dae" / "output.json"),
        "plan_path": str(doc_dir / "plan" / "primary.json"),
        "ocr_path": str(doc_dir / "preprocessing" / "ocr" / "text.json"),
        "structuring_extracted_path": str(doc_dir / "structuring" / "extracted.json"),
        "structuring_sections_path": str(doc_dir / "structuring" / "sections.json"),
        "structuring_summary_path": str(doc_dir / "structuring" / "summaries.json"),
        "validation_results_path": str(doc_dir / "validation" / "results.json"),
        "validation_normalized_path": str(doc_dir / "validation" / "normalized.json"),
        "recovery_attempts_path": str(doc_dir / "recovery" / "attempts.jsonl"),
        "recovery_escalations_path": str(doc_dir / "recovery" / "escalations.jsonl"),
        "chunking_path": str(doc_dir / "chunking" / "chunks.jsonl"),
        "embeddings_index_path": str(doc_dir / "embeddings" / "index"),
        "provenance_graph_path": str(doc_dir / "provenance" / "graph.json"),
        "provenance_events_path": str(doc_dir / "provenance" / "events.jsonl"),
        "transparency_path": str(doc_dir / "provenance" / "transparency.json"),
        "artifacts_path": str(doc_dir / "provenance" / "artifacts"),
    }

    field_map = _map_entities(result.entities)
    for field in expected_fields:
        key = field["key"]
        status = requested_field_status.get(key, {})
        value = str(status.get("value", "")).strip()
        confidence = ""
        if status:
            raw_conf = status.get("confidence")
            if isinstance(raw_conf, (int, float)):
                confidence = f"{float(raw_conf):.3f}"
            elif raw_conf is not None:
                confidence = str(raw_conf)
        page = ""
        span = ""
        source = str(status.get("source", "")).strip() if status else ""

        if not value:
            value, confidence, page, span, source = field_map.get(key, ("", "", "", "", ""))

        record[f"field_{key}"] = value
        record[f"field_{key}_confidence"] = confidence
        record[f"field_{key}_page"] = page
        record[f"field_{key}_span"] = span
        record[f"field_{key}_source"] = source

    if expected_keys:
        record["structured_fields"] = {key: record.get(f"field_{key}", "") for key in expected_keys}
        record["missing_fields"] = [
            key
            for key in expected_keys
            if str((requested_field_status.get(key) or {}).get("status", "")).strip() == "missing"
            or not str(record.get(f"field_{key}", "")).strip()
        ]
    else:
        record["structured_fields"] = dict(entities_dict)
        record["missing_fields"] = []

    return record

def _map_entities(entities) -> dict[str, tuple[str, str, str, str, str]]:
    mapped: dict[str, tuple[str, str, str, str, str]] = {}
    for entity in entities:
        key = _normalize_key(entity.key)
        value = str(entity.value)
        confidence = f"{entity.confidence:.3f}" if entity.confidence is not None else ""
        page = str(entity.page) if entity.page is not None else ""
        span = ""
        if entity.span:
            span = f"{entity.span[0]}:{entity.span[1]}"
        source = ""
        if entity.metadata:
            source = str(entity.metadata.get("source") or entity.metadata.get("explicit") or "")
        existing = mapped.get(key)
        if existing:
            value = f"{existing[0]} | {value}"
        mapped[key] = (value, confidence, page, span, source)
    return mapped

def _find_original_path(doc_dir: Path) -> str:
    input_dir = doc_dir / "input"
    matches = list(input_dir.glob("original.*"))
    if matches:
        return str(matches[0])
    return str(input_dir / "original")

def _write_docs_excel(run_dir: Path, records: list[dict[str, object]]) -> None:
    if not records:
        return
    headers = _collect_headers(records)
    try:
        import openpyxl
    except Exception:
        csv_path = run_dir / "docs.csv"
        with csv_path.open("w", encoding="utf-8") as handle:
            handle.write(",".join(headers) + "\n")
            for record in records:
                row = [str(record.get(header, "")).replace(",", " ") for header in headers]
                handle.write(",".join(row) + "\n")
        return
    path = run_dir / "docs.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "docs"
    sheet.append(headers)
    for record in records:
        sheet.append([_excel_safe(record.get(header, "")) for header in headers])
    wb.save(path)

def _collect_headers(records: list[dict[str, object]]) -> list[str]:
    if not records:
        return []
    keys = set()
    for record in records:
        keys.update(record.keys())
    ordered = [key for key in BASE_DOC_HEADERS if key in keys]
    remaining = sorted(key for key in keys if key not in ordered)
    return ordered + remaining

def _entity_records(records: list[dict[str, object]], expected_fields: list[dict[str, str]]) -> list[dict[str, object]]:
    if not expected_fields:
        return []
    field_keys = [field["key"] for field in expected_fields]
    entity_records = []
    for record in records:
        row = {"doc_id": record.get("doc_id", "")}
        for key in field_keys:
            row[key] = record.get(f"field_{key}", "")
        entity_records.append(row)
    return entity_records

def _write_entities_excel(
    run_dir: Path,
    records: list[dict[str, object]],
    expected_fields: list[dict[str, str]],
) -> None:
    entity_records = _entity_records(records, expected_fields)
    if not entity_records:
        return
    headers = ["doc_id"] + [field["key"] for field in expected_fields]
    try:
        import openpyxl
    except Exception:
        csv_path = run_dir / "entities.csv"
        with csv_path.open("w", encoding="utf-8") as handle:
            handle.write(",".join(headers) + "\n")
            for record in entity_records:
                row = [str(record.get(header, "")).replace(",", " ") for header in headers]
                handle.write(",".join(row) + "\n")
        return
    path = run_dir / "entities.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "entities"
    sheet.append(headers)
    for record in entity_records:
        sheet.append([record.get(header, "") for header in headers])
    wb.save(path)

def _write_run_excel(run_dir: Path, run_payload: dict[str, object], metrics_payload: dict[str, object]) -> None:
    try:
        import openpyxl
    except Exception:
        csv_path = run_dir / "run.csv"
        payload = {**run_payload, **metrics_payload}
        headers = sorted(payload.keys())
        with csv_path.open("w", encoding="utf-8") as handle:
            handle.write(",".join(headers) + "\n")
            handle.write(",".join(str(payload.get(header, "")) for header in headers))
        return
    path = run_dir / "run.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "run"
    payload = {**run_payload, **metrics_payload}
    headers = sorted(payload.keys())
    sheet.append(headers)
    sheet.append([_excel_safe(payload.get(header, "")) for header in headers])
    wb.save(path)

def _persist_run_telemetry(
    run_dir: Path,
    run_id: str,
    doc_records: list[dict[str, object]],
    metrics_payload: dict[str, object],
    preflight_warnings: list[str],
) -> dict[str, object]:
    dsn = (os.getenv("POSTGRES_URI") or "").strip()
    if not dsn:
        return {"enabled": False, "status": "skipped", "reason": "POSTGRES_URI_not_set"}
    try:
        import psycopg
    except Exception as exc:
        return {"enabled": False, "status": "skipped", "warning": f"psycopg_unavailable: {exc}"}

    try:
        with psycopg.connect(dsn, autocommit=True) as conn:
            with conn.cursor() as cur:
                cur.execute(
                    """
                    CREATE TABLE IF NOT EXISTS idpo_run_telemetry (
                        run_id TEXT PRIMARY KEY,
                        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                        doc_count INTEGER NOT NULL,
                        processed_doc_count INTEGER NOT NULL,
                        failed_count INTEGER NOT NULL,
                        duration_ms DOUBLE PRECISION NOT NULL,
                        preflight_warnings TEXT NOT NULL,
                        metrics_json JSONB NOT NULL
                    )
                    """
                )
                cur.execute(
                    """
                    CREATE TABLE IF NOT EXISTS idpo_doc_telemetry (
                        run_id TEXT NOT NULL,
                        doc_id TEXT NOT NULL,
                        created_at TIMESTAMPTZ NOT NULL DEFAULT NOW(),
                        model_mode TEXT,
                        model_used TEXT,
                        ocr_confidence DOUBLE PRECISION,
                        cloud_total_tokens INTEGER,
                        cloud_estimated_cost_usd DOUBLE PRECISION,
                        fallback_path TEXT,
                        extraction_success BOOLEAN,
                        unresolved_fields TEXT[] NOT NULL DEFAULT '{}',
                        latency_ms DOUBLE PRECISION,
                        llm_diag_json JSONB NOT NULL,
                        PRIMARY KEY (run_id, doc_id)
                    )
                    """
                )
                cur.execute(
                    """
                    INSERT INTO idpo_run_telemetry
                        (run_id, doc_count, processed_doc_count, failed_count, duration_ms, preflight_warnings, metrics_json)
                    VALUES (%s, %s, %s, %s, %s, %s, %s::jsonb)
                    ON CONFLICT (run_id) DO UPDATE SET
                        doc_count = EXCLUDED.doc_count,
                        processed_doc_count = EXCLUDED.processed_doc_count,
                        failed_count = EXCLUDED.failed_count,
                        duration_ms = EXCLUDED.duration_ms,
                        preflight_warnings = EXCLUDED.preflight_warnings,
                        metrics_json = EXCLUDED.metrics_json
                    """,
                    (
                        run_id,
                        int(metrics_payload.get("doc_count", 0) or 0),
                        int(metrics_payload.get("processed_doc_count", 0) or 0),
                        int(metrics_payload.get("failed_count", 0) or 0),
                        float(metrics_payload.get("duration_ms", 0.0) or 0.0),
                        " | ".join(preflight_warnings),
                        json.dumps(metrics_payload),
                    ),
                )

                persisted_docs = 0
                for record in doc_records:
                    doc_id = str(record.get("doc_id", "")).strip()
                    if not doc_id:
                        continue
                    llm_diag = record.get("llm_diagnostics", {})
                    if not isinstance(llm_diag, dict):
                        llm_diag = {}
                    cloud_usage = llm_diag.get("cloud_usage", {})
                    if not isinstance(cloud_usage, dict):
                        cloud_usage = {}

                    ocr_conf = _read_ocr_confidence(record.get("ocr_path"))
                    fallback_path = _fallback_path(llm_diag)
                    missing_fields = record.get("missing_fields", [])
                    if not isinstance(missing_fields, list):
                        missing_fields = []
                    extraction_success = bool(record.get("entity_count", 0)) and not bool(missing_fields)

                    cur.execute(
                        """
                        INSERT INTO idpo_doc_telemetry
                            (run_id, doc_id, model_mode, model_used, ocr_confidence, cloud_total_tokens,
                             cloud_estimated_cost_usd, fallback_path, extraction_success, unresolved_fields,
                             latency_ms, llm_diag_json)
                        VALUES (%s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s, %s::jsonb)
                        ON CONFLICT (run_id, doc_id) DO UPDATE SET
                            model_mode = EXCLUDED.model_mode,
                            model_used = EXCLUDED.model_used,
                            ocr_confidence = EXCLUDED.ocr_confidence,
                            cloud_total_tokens = EXCLUDED.cloud_total_tokens,
                            cloud_estimated_cost_usd = EXCLUDED.cloud_estimated_cost_usd,
                            fallback_path = EXCLUDED.fallback_path,
                            extraction_success = EXCLUDED.extraction_success,
                            unresolved_fields = EXCLUDED.unresolved_fields,
                            latency_ms = EXCLUDED.latency_ms,
                            llm_diag_json = EXCLUDED.llm_diag_json
                        """,
                        (
                            run_id,
                            doc_id,
                            str(record.get("model_mode", "") or ""),
                            str(record.get("model_used", "") or ""),
                            ocr_conf,
                            int(cloud_usage.get("total_tokens", 0) or 0),
                            float(cloud_usage.get("estimated_cost_usd", 0.0) or 0.0),
                            fallback_path,
                            extraction_success,
                            [str(item) for item in missing_fields],
                            float(record.get("latency_ms", 0.0) or 0.0),
                            json.dumps(llm_diag),
                        ),
                    )
                    persisted_docs += 1

        return {"enabled": True, "status": "ok", "persisted_docs": persisted_docs}
    except Exception as exc:
        return {"enabled": False, "status": "error", "warning": f"telemetry_db_write_failed: {exc}"}

def _fallback_path(llm_diag: dict[str, object]) -> str:
    local_attempted = bool(llm_diag.get("local_attempted"))
    cloud_attempted = bool(llm_diag.get("cloud_attempted"))
    flagship_attempted = bool(llm_diag.get("flagship_attempted"))
    cloud_succeeded = bool(llm_diag.get("cloud_succeeded"))
    local_succeeded = bool(llm_diag.get("local_succeeded"))
    requested_mode = str(llm_diag.get("requested_mode", "") or "").strip().lower()
    if flagship_attempted:
        return "local_or_cloud->flagship" if (local_attempted or cloud_attempted) else "flagship_only"
    if local_attempted and cloud_attempted:
        if requested_mode == "cloud" and local_succeeded and not cloud_succeeded:
            return "cloud->local"
        return "local->cloud"
    if cloud_attempted and cloud_succeeded:
        return "cloud"
    if local_attempted and local_succeeded:
        return "local"
    return "none"

def _read_ocr_confidence(path_like: object) -> float:
    if not path_like:
        return 0.0
    try:
        path = Path(str(path_like))
        if not path.exists():
            return 0.0
        payload = json.loads(path.read_text(encoding="utf-8"))
        value = payload.get("confidence", 0.0)
        return float(value or 0.0)
    except Exception:
        return 0.0

def _excel_safe(value: object) -> object:
    if isinstance(value, (list, dict)):
        return json.dumps(value)
    return value
