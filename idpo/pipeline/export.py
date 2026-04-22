"""Result export to CSV and XLSX formats."""

from __future__ import annotations

import csv
import json
import shutil
from pathlib import Path
from typing import Any, Iterable, TYPE_CHECKING

if TYPE_CHECKING:
    from idpo.core.types import Document
    from idpo.pipeline.results import PipelineResult

def export_structured_results(results: list[PipelineResult], output_path: str) -> dict[str, Any]:
    path = Path(output_path)
    try:
        import openpyxl
    except Exception:
        return _export_csv(results, path)
    return _export_xlsx(results, path, openpyxl)

def _export_xlsx(results: list[PipelineResult], path: Path, openpyxl) -> dict[str, Any]:
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = openpyxl.Workbook()
    explicit_sheet = wb.active
    explicit_sheet.title = "explicit_entities"
    default_sheet = wb.create_sheet("default_entities")
    meta_sheet = wb.create_sheet("metadata")

    explicit_keys = _collect_explicit_keys(results)
    explicit_headers = ["doc_id", "policy_id", "policy_name", "model_used", "explicit_source"] + explicit_keys
    explicit_sheet.append(explicit_headers)
    for result in results:
        explicit_values = _entities_by_key(result.explicit_entities)
        row = [
            result.doc_id,
            result.policy_id,
            result.policy_name,
            result.model_used,
            result.explicit_source,
        ] + [explicit_values.get(key, "") for key in explicit_keys]
        explicit_sheet.append(row)

    default_headers = ["doc_id", "entity_key", "entity_value", "confidence", "page", "metadata"]
    default_sheet.append(default_headers)
    for result in results:
        for entity in result.default_entities:
            default_sheet.append(
                [
                    result.doc_id,
                    entity.key,
                    entity.value,
                    entity.confidence,
                    entity.page or "",
                    _json_cell(entity.metadata),
                ]
            )

    meta_headers = [
        "doc_id",
        "path",
        "source",
        "doc_type",
        "category",
        "bin",
        "profile",
        "vendor_cluster",
        "quality_bin",
        "complexity_bin",
        "content_profile",
        "quality_score",
        "complexity_score",
        "policy_id",
        "policy_name",
        "policy_reason",
        "plan_id",
        "plan_source",
        "plan_budgets",
        "plan_retrieval",
        "plan_token_budget",
        "model_used",
        "model_mode",
        "model_reason",
        "preprocess_route",
        "preprocess_artifacts",
        "explicit_source",
        "entity_count",
        "explicit_count",
        "default_count",
        "validation_issue_count",
        "recovered_fields",
        "field_details",
        "llm_diagnostics",
        "escalated",
        "queue_partition",
        "queue_wait_ms",
        "latency_ms",
    ]
    meta_sheet.append(meta_headers)
    for result in results:
        meta_sheet.append(
            [
                result.doc_id,
                result.path,
                result.source,
                result.doc_type,
                result.category,
                result.bin,
                result.profile,
                result.vendor_cluster,
                result.quality_bin,
                result.complexity_bin,
                result.content_profile,
                result.quality_score,
                result.complexity_score,
                result.policy_id,
                result.policy_name,
                result.policy_reason,
                _json_cell(result.plan_id),
                result.plan_source,
                _json_cell(result.plan_budgets),
                _json_cell(result.plan_retrieval),
                _json_cell(result.plan_token_budget),
                result.model_used,
                result.model_mode,
                result.model_reason,
                ",".join(result.preprocess_route),
                ",".join(result.preprocess_artifacts),
                result.explicit_source,
                len(result.entities),
                len(result.explicit_entities),
                len(result.default_entities),
                len(result.validation_issues),
                _json_cell(result.recovered_fields),
                _json_cell(getattr(result, "field_details", {})),
                _json_cell(getattr(result, "llm_diagnostics", {})),
                result.escalated,
                result.queue_partition,
                result.queue_wait_ms,
                result.latency_ms,
            ]
        )

    wb.save(path)
    return {"format": "xlsx", "path": str(path), "sheets": ["explicit_entities", "default_entities", "metadata"]}

def _export_csv(results: list[PipelineResult], path: Path) -> dict[str, Any]:
    path.parent.mkdir(parents=True, exist_ok=True)
    base = path.with_suffix("")
    explicit_path = base.with_name(f"{base.name}_explicit.csv")
    default_path = base.with_name(f"{base.name}_default.csv")
    meta_path = base.with_name(f"{base.name}_metadata.csv")

    explicit_keys = _collect_explicit_keys(results)
    explicit_headers = ["doc_id", "policy_id", "policy_name", "model_used", "explicit_source"] + explicit_keys
    with explicit_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(explicit_headers)
        for result in results:
            explicit_values = _entities_by_key(result.explicit_entities)
            row = [
                result.doc_id,
                result.policy_id,
                result.policy_name,
                result.model_used,
                result.explicit_source,
            ] + [explicit_values.get(key, "") for key in explicit_keys]
            writer.writerow(row)

    default_headers = ["doc_id", "entity_key", "entity_value", "confidence", "page", "metadata"]
    with default_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(default_headers)
        for result in results:
            for entity in result.default_entities:
                writer.writerow(
                    [
                        result.doc_id,
                        entity.key,
                        entity.value,
                        entity.confidence,
                        entity.page or "",
                        _json_cell(entity.metadata),
                    ]
                )

    meta_headers = [
        "doc_id",
        "path",
        "source",
        "doc_type",
        "category",
        "bin",
        "profile",
        "vendor_cluster",
        "quality_bin",
        "complexity_bin",
        "content_profile",
        "quality_score",
        "complexity_score",
        "policy_id",
        "policy_name",
        "policy_reason",
        "plan_id",
        "plan_source",
        "plan_budgets",
        "plan_retrieval",
        "plan_token_budget",
        "model_used",
        "model_mode",
        "model_reason",
        "preprocess_route",
        "preprocess_artifacts",
        "explicit_source",
        "entity_count",
        "explicit_count",
        "default_count",
        "validation_issue_count",
        "recovered_fields",
        "field_details",
        "llm_diagnostics",
        "escalated",
        "queue_partition",
        "queue_wait_ms",
        "latency_ms",
    ]
    with meta_path.open("w", newline="", encoding="utf-8") as handle:
        writer = csv.writer(handle)
        writer.writerow(meta_headers)
        for result in results:
            writer.writerow(
                [
                    result.doc_id,
                    result.path,
                    result.source,
                    result.doc_type,
                    result.category,
                    result.bin,
                    result.profile,
                    result.vendor_cluster,
                    result.quality_bin,
                    result.complexity_bin,
                    result.content_profile,
                    result.quality_score,
                    result.complexity_score,
                    result.policy_id,
                    result.policy_name,
                    result.policy_reason,
                    _json_cell(result.plan_id),
                    result.plan_source,
                    _json_cell(result.plan_budgets),
                    _json_cell(result.plan_retrieval),
                    _json_cell(result.plan_token_budget),
                    result.model_used,
                    result.model_mode,
                    result.model_reason,
                    ",".join(result.preprocess_route),
                    ",".join(result.preprocess_artifacts),
                    result.explicit_source,
                    len(result.entities),
                    len(result.explicit_entities),
                    len(result.default_entities),
                    len(result.validation_issues),
                    _json_cell(result.recovered_fields),
                    _json_cell(getattr(result, "field_details", {})),
                    _json_cell(getattr(result, "llm_diagnostics", {})),
                    result.escalated,
                    result.queue_partition,
                    result.queue_wait_ms,
                    result.latency_ms,
                ]
            )
    return {"format": "csv", "paths": [str(explicit_path), str(default_path), str(meta_path)]}

def _collect_explicit_keys(results: Iterable[PipelineResult]) -> list[str]:
    keys = set()
    for result in results:
        for entity in result.explicit_entities:
            keys.add(entity.key)
    return sorted(keys)

def _entities_by_key(entities) -> dict[str, str]:
    values: dict[str, tuple[float, str]] = {}
    for entity in entities:
        current = values.get(entity.key)
        if not current or entity.confidence > current[0]:
            values[entity.key] = (entity.confidence, entity.value)
    return {key: value for key, (_, value) in values.items()}

def _json_cell(value: Any) -> str:
    if value is None:
        return ""
    return json.dumps(value, sort_keys=True)

def export_simple_workspace(
    docs: list["Document"],
    results: list["PipelineResult"],
    output_dir: str,
) -> dict[str, Any]:
    root = Path(output_dir)
    raw_root = root / "raw_categorized"
    text_root = root / "text_extracted"
    raw_root.mkdir(parents=True, exist_ok=True)
    text_root.mkdir(parents=True, exist_ok=True)

    result_map = {result.doc_id: result for result in results}
    manifest = []
    copied_inputs = set()

    for doc in docs:
        result = result_map.get(doc.doc_id)
        category = result.category if result else "unknown"
        category_dir = raw_root / category
        category_dir.mkdir(parents=True, exist_ok=True)

        input_path = Path(doc.path)
        raw_path = None
        if input_path.exists() and input_path.is_file() and input_path.suffix.lower() not in {".json"}:
            raw_path = category_dir / f"{doc.doc_id}{input_path.suffix.lower()}"
            if input_path not in copied_inputs:
                shutil.copy2(input_path, raw_path)
                copied_inputs.add(input_path)
        else:
            raw_path = category_dir / f"{doc.doc_id}_raw.txt"
            raw_text = "\n".join((page.raw_text or page.text) for page in doc.pages)
            raw_path.write_text(raw_text, encoding="utf-8")

        extracted_path = text_root / f"{doc.doc_id}.txt"
        extracted_text = "\n".join(page.text for page in doc.pages)
        extracted_path.write_text(extracted_text, encoding="utf-8")

        manifest.append(
            {
                "doc_id": doc.doc_id,
                "category": category,
                "input_path": str(input_path),
                "raw_path": str(raw_path),
                "text_extracted_path": str(extracted_path),
            }
        )

    manifest_path = root / "simple_manifest.json"
    manifest_path.write_text(json.dumps({"documents": manifest}, indent=2), encoding="utf-8")
    return {"output_dir": str(root), "manifest": str(manifest_path), "count": len(manifest)}
