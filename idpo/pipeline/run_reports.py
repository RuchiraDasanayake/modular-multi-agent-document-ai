"""Run-level reporting helpers for inbox processing.

This module owns report-shaped outputs derived from completed document records.
It is intentionally separate from the main runtime entrypoint so reporting
concerns do not keep expanding the core inbox orchestration path.
"""

from __future__ import annotations

import json
from pathlib import Path

BASE_DOC_HEADERS = [
    "doc_id",
    "source",
    "doc_type",
    "vendor_cluster",
    "semantic_doc_type",
    "layout_cluster_id",
    "route_bin",
    "route_specialization",
    "quality_bin",
    "complexity_bin",
    "content_profile",
    "quality_score",
    "complexity_score",
    "analysis_confidence",
    "analysis_confidence_type",
    "policy_id",
    "policy_name",
    "static_policy_id",
    "static_policy_name",
    "static_policy_reason",
    "advice_source",
    "advice_status",
    "advice_accepted",
    "policy_changed",
    "fallback_reason",
    "advisor_recommendation",
    "fallback_candidates",
    "operator_stack",
    "defer_status",
    "defer_reason",
    "plan_id",
    "plan_source",
    "model_used",
    "model_mode",
    "model_reason",
    "latency_ms",
    "preprocess_route",
    "ocr_engine_used",
    "ocr_avg_confidence",
    "ocr_rerun_attempted_pages",
    "ocr_rerun_used_pages",
    "fallback_path",
    "unresolved_field_count",
    "schema_required_field_count",
    "schema_extracted_required_count",
    "schema_missing_required_count",
    "schema_hallucinated_field_count",
    "schema_grounded_field_count",
    "schema_coverage",
    "schema_grounding_coverage",
    "schema_valid",
    "entity_count",
    "explicit_entity_count",
    "default_entity_count",
    "validation_issue_count",
    "recovered_fields_count",
    "reward",
    "reward_type",
    "escalation_status",
    "requested_field_status",
    "llm_diagnostics",
    "recovery_log",
    "escalation_info",
    "escalated",
    "chunk_count",
    "summary",
    "input_original_path",
    "input_manifest_path",
    "input_hashes_path",
    "dae_output_path",
    "plan_path",
    "ocr_path",
    "structuring_extracted_path",
    "structuring_sections_path",
    "structuring_summary_path",
    "validation_results_path",
    "validation_normalized_path",
    "recovery_attempts_path",
    "recovery_escalations_path",
    "decision_chain_path",
    "reward_path",
    "recovery_trace_path",
    "chunking_path",
    "embeddings_index_path",
    "provenance_graph_path",
    "provenance_events_path",
    "transparency_path",
    "artifacts_path",
]


def aggregate_reward_metrics(records: list[dict[str, object]]) -> dict[str, object]:
    if not records:
        return {
            "avg_reward": 0.0,
            "direct_reward_count": 0,
            "proxy_reward_count": 0,
            "avg_direct_f1": 0.0,
            "avg_schema_coverage": 0.0,
            "avg_schema_grounding_coverage": 0.0,
            "schema_valid_rate": 0.0,
            "defer_requested_count": 0,
            "defer_blocked_count": 0,
            "advice_accepted_count": 0,
            "policy_changed_count": 0,
        }
    rewards = [float(record.get("reward", 0.0) or 0.0) for record in records]
    direct_reward_types = {"direct_set_based", "policy_declared", "canonical_fallback"}
    direct_scores = [
        float(((record.get("reward_details") or {}).get("f1", 0.0)) or 0.0)
        for record in records
        if str(record.get("reward_type", "")).strip() in direct_reward_types
    ]
    direct_count = sum(1 for record in records if str(record.get("reward_type", "")).strip() in direct_reward_types)
    proxy_count = sum(1 for record in records if str(record.get("reward_type", "")).strip() == "proxy")
    schema_coverages = [float(record.get("schema_coverage", 0.0) or 0.0) for record in records]
    grounding_coverages = [float(record.get("schema_grounding_coverage", 0.0) or 0.0) for record in records]
    schema_valid_count = sum(1 for record in records if bool(record.get("schema_valid", False)))
    defer_requested_count = sum(1 for record in records if str(record.get("defer_status", "")).strip() == "requested")
    defer_blocked_count = sum(1 for record in records if str(record.get("defer_status", "")).strip() == "blocked")
    advice_accepted_count = sum(1 for record in records if bool(record.get("advice_accepted", False)))
    policy_changed_count = sum(1 for record in records if bool(record.get("policy_changed", False)))
    return {
        "avg_reward": round(sum(rewards) / len(rewards), 4),
        "direct_reward_count": direct_count,
        "proxy_reward_count": proxy_count,
        "avg_direct_f1": round(sum(direct_scores) / len(direct_scores), 4) if direct_scores else 0.0,
        "avg_schema_coverage": round(sum(schema_coverages) / len(schema_coverages), 4) if schema_coverages else 0.0,
        "avg_schema_grounding_coverage": round(sum(grounding_coverages) / len(grounding_coverages), 4) if grounding_coverages else 0.0,
        "schema_valid_rate": round(schema_valid_count / len(records), 4),
        "defer_requested_count": defer_requested_count,
        "defer_blocked_count": defer_blocked_count,
        "advice_accepted_count": advice_accepted_count,
        "policy_changed_count": policy_changed_count,
    }


def write_decision_chain_summary(run_dir: Path, records: list[dict[str, object]]) -> None:
    summary = []
    for record in records:
        summary.append(
            {
                "doc_id": record.get("doc_id", ""),
                "filename": record.get("filename", ""),
                "policy_id": record.get("policy_id", ""),
                "semantic_doc_type": record.get("semantic_doc_type", "unknown"),
                "static_policy_id": record.get("static_policy_id"),
                "static_policy_name": record.get("static_policy_name"),
                "static_policy_reason": record.get("static_policy_reason"),
                "advice_source": record.get("advice_source", ""),
                "advice_status": record.get("advice_status", "static"),
                "advice_id": record.get("advice_id"),
                "advice_confidence": record.get("advice_confidence", 0.0),
                "retrieval_source": record.get("retrieval_source"),
                "rabi_used": record.get("rabi_used", False),
                "advice_accepted": record.get("advice_accepted", False),
                "policy_changed": record.get("policy_changed", False),
                "advisor_recommendation": record.get("advisor_recommendation"),
                "fallback_candidates": record.get("fallback_candidates", []),
                "fallback_reason": record.get("fallback_reason"),
                "blocked_policy": record.get("blocked_policy", False),
                "rollback_triggered": record.get("rollback_triggered", False),
                "rollback_reason": record.get("rollback_reason"),
                "defer_status": record.get("defer_status", "not_needed"),
                "defer_reason": record.get("defer_reason"),
                "route_bin": record.get("route_bin", ""),
                "route_specialization": record.get("route_specialization"),
                "operator_stack": record.get("operator_stack", []),
                "schema_coverage": record.get("schema_coverage", 0.0),
                "schema_grounding_coverage": record.get("schema_grounding_coverage", 0.0),
                "schema_valid": record.get("schema_valid", False),
                "schema_missing_required_count": record.get("schema_missing_required_count", 0),
                "schema_hallucinated_field_count": record.get("schema_hallucinated_field_count", 0),
                "reward": record.get("reward", 0.0),
                "reward_type": record.get("reward_type", ""),
                "spent_budget_units": record.get("spent_budget_units", 0.0),
                "feedback_logged": record.get("feedback_logged", False),
                "feedback_error": record.get("feedback_error"),
                "validation_issue_count": record.get("validation_issue_count", 0),
                "recovered_fields_count": record.get("recovered_fields_count", 0),
                "escalation_status": record.get("escalation_status", "not_needed"),
                "latency_ms": record.get("latency_ms", 0.0),
            }
        )
    (run_dir / "decision_chain_summary.json").write_text(json.dumps(summary, indent=2), encoding="utf-8")


def write_docs_excel(run_dir: Path, records: list[dict[str, object]]) -> None:
    if not records:
        return
    headers = _collect_headers(records)
    try:
        import openpyxl
    except Exception:
        csv_path = run_dir / "docs.csv"
        with csv_path.open("w", encoding="utf-8") as handle:
            handle.write(",".join(headers) + "\n")
            for record in records:
                row = [str(record.get(header, "")).replace(",", " ") for header in headers]
                handle.write(",".join(row) + "\n")
        return
    path = run_dir / "docs.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "docs"
    sheet.append(headers)
    for record in records:
        sheet.append([_excel_safe(record.get(header, "")) for header in headers])
    wb.save(path)


def entity_records(records: list[dict[str, object]], expected_fields: list[dict[str, str]]) -> list[dict[str, object]]:
    if not expected_fields:
        return []
    field_keys = [field["key"] for field in expected_fields]
    rows = []
    for record in records:
        row = {"doc_id": record.get("doc_id", "")}
        for key in field_keys:
            row[key] = record.get(f"field_{key}", "")
        rows.append(row)
    return rows


def write_entities_excel(
    run_dir: Path,
    records: list[dict[str, object]],
    expected_fields: list[dict[str, str]],
) -> None:
    rows = entity_records(records, expected_fields)
    if not rows:
        return
    headers = ["doc_id"] + [field["key"] for field in expected_fields]
    try:
        import openpyxl
    except Exception:
        csv_path = run_dir / "entities.csv"
        with csv_path.open("w", encoding="utf-8") as handle:
            handle.write(",".join(headers) + "\n")
            for record in rows:
                row = [str(record.get(header, "")).replace(",", " ") for header in headers]
                handle.write(",".join(row) + "\n")
        return
    path = run_dir / "entities.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "entities"
    sheet.append(headers)
    for record in rows:
        sheet.append([record.get(header, "") for header in headers])
    wb.save(path)


def write_run_excel(run_dir: Path, run_payload: dict[str, object], metrics_payload: dict[str, object]) -> None:
    try:
        import openpyxl
    except Exception:
        csv_path = run_dir / "run.csv"
        payload = {**run_payload, **metrics_payload}
        headers = sorted(payload.keys())
        with csv_path.open("w", encoding="utf-8") as handle:
            handle.write(",".join(headers) + "\n")
            handle.write(",".join(str(payload.get(header, "")) for header in headers))
        return
    path = run_dir / "run.xlsx"
    wb = openpyxl.Workbook()
    sheet = wb.active
    sheet.title = "run"
    payload = {**run_payload, **metrics_payload}
    headers = sorted(payload.keys())
    sheet.append(headers)
    sheet.append([_excel_safe(payload.get(header, "")) for header in headers])
    wb.save(path)


def _collect_headers(records: list[dict[str, object]]) -> list[str]:
    if not records:
        return []
    keys = set()
    for record in records:
        keys.update(record.keys())
    ordered = [key for key in BASE_DOC_HEADERS if key in keys]
    remaining = sorted(key for key in keys if key not in ordered)
    return ordered + remaining


def _excel_safe(value: object) -> object:
    if isinstance(value, (list, dict)):
        return json.dumps(value)
    return value
